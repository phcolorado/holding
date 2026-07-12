import os
from calendar import monthrange
from datetime import date, timedelta
from decimal import Decimal

from django.core.exceptions import ValidationError
from django.db import models
from django.test import TestCase, TransactionTestCase, Client, skipUnlessDBFeature
from django.contrib.auth.models import User
from core.test_utils import com_leitura
from django.urls import reverse
from django.utils import timezone

from patrimonio.models import Imovel, Pessoa, Contrato, EncargoContrato
from financeiro.models import (
    ReceitaAluguel, ReceitaAluguelItem, RecebimentoReceita, Despesa, receitas_inadimplentes_qs,
    valor_total_devido_expr,
)
from financeiro.services import gerar_receitas_para_contrato, gerar_receitas_mes, contratos_para_geracao_mes


def _criar_base():
    """Cria imóvel e locatário reutilizáveis entre testes."""
    imovel = Imovel.objects.create(
        nome='Apartamento Teste',
        endereco='Rua A, 100',
        cidade='São Paulo',
        estado='SP',
    )
    locatario = Pessoa.objects.create(nome='Locatário Teste', tipo='locatario')
    return imovel, locatario


def _criar_contrato(imovel, locatario, data_inicio, data_fim, status='ativo', dia_vencimento=10, **kwargs):
    defaults = dict(valor_aluguel=Decimal('2000.00'))
    defaults.update(kwargs)
    return Contrato.objects.create(
        imovel=imovel,
        locatario=locatario,
        data_inicio=data_inicio,
        data_fim=data_fim,
        dia_vencimento=dia_vencimento,
        status=status,
        **defaults,
    )


def _dar_permissoes(user, *codenames):
    """Concede as permissões (por codename) exigidas pelas ações de escrita das views."""
    from django.contrib.auth.models import Permission
    user.user_permissions.add(*Permission.objects.filter(codename__in=codenames))


class GerarReceitasParaContratoTest(TestCase):
    """Testa a função gerar_receitas_para_contrato."""

    def setUp(self):
        self.imovel, self.locatario = _criar_base()
        # Contrato de 3 meses: jan–mar 2024
        self.contrato = _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(2024, 1, 1),
            data_fim=date(2024, 3, 31),
        )

    def test_contrato_ativo_gera_receitas_mensais(self):
        """Contrato ativo de 3 meses gera exatamente 3 receitas."""
        criadas, existiam = gerar_receitas_para_contrato(self.contrato)

        self.assertEqual(criadas, 3)
        self.assertEqual(existiam, 0)
        self.assertEqual(ReceitaAluguel.objects.count(), 3)

        meses = list(
            ReceitaAluguel.objects.order_by('competencia_mes')
            .values_list('competencia_mes', flat=True)
        )
        self.assertEqual(meses, [1, 2, 3])

    def test_contrato_inativo_nao_gera_receitas(self):
        """Contratos com status diferente de 'ativo' são ignorados."""
        for status in ('encerrado', 'rescindido', 'suspenso'):
            self.contrato.status = status
            self.contrato.save()

            criadas, existiam = gerar_receitas_para_contrato(self.contrato)

            self.assertEqual(criadas, 0, msg=f'Status {status} não deveria gerar receitas')
            self.assertEqual(existiam, 0)

        self.assertEqual(ReceitaAluguel.objects.count(), 0)

    def test_nao_cria_duplicidade(self):
        """Segunda chamada não cria registros duplicados."""
        criadas1, existiam1 = gerar_receitas_para_contrato(self.contrato)
        criadas2, existiam2 = gerar_receitas_para_contrato(self.contrato)

        self.assertEqual(criadas1, 3)
        self.assertEqual(existiam1, 0)
        self.assertEqual(criadas2, 0)
        self.assertEqual(existiam2, 3)
        # Sem duplicatas no banco
        self.assertEqual(ReceitaAluguel.objects.count(), 3)

    def test_receita_gerada_usa_imovel_do_contrato(self):
        """O imóvel da receita gerada é sempre o imóvel do contrato."""
        gerar_receitas_para_contrato(self.contrato)

        for receita in ReceitaAluguel.objects.all():
            self.assertEqual(
                receita.imovel_id, self.contrato.imovel_id,
                msg='Imóvel da receita difere do imóvel do contrato',
            )

    def test_valores_e_status_corretos(self):
        """Receita gerada tem valor_previsto, data_vencimento e status corretos."""
        gerar_receitas_para_contrato(self.contrato)

        receita_jan = ReceitaAluguel.objects.get(competencia_mes=1, competencia_ano=2024)
        self.assertEqual(receita_jan.data_vencimento, date(2024, 1, 10))
        self.assertEqual(receita_jan.valor_previsto, Decimal('2000.00'))
        self.assertEqual(receita_jan.status, 'previsto')

    def test_dia_vencimento_ajustado_para_fim_do_mes(self):
        """Dia 31 em fevereiro (ano bissexto 2024) deve cair no dia 29."""
        self.contrato.dia_vencimento = 31
        self.contrato.save()

        gerar_receitas_para_contrato(self.contrato)

        receita_fev = ReceitaAluguel.objects.get(competencia_mes=2, competencia_ano=2024)
        self.assertEqual(receita_fev.data_vencimento, date(2024, 2, 29))

    def test_gerar_apenas_periodo_especifico(self):
        """Com data_inicio/data_fim informados, gera somente receitas do intervalo."""
        criadas, _ = gerar_receitas_para_contrato(
            self.contrato,
            data_inicio=date(2024, 2, 1),
            data_fim=date(2024, 2, 29),
        )

        self.assertEqual(criadas, 1)
        self.assertTrue(
            ReceitaAluguel.objects.filter(competencia_mes=2, competencia_ano=2024).exists()
        )
        self.assertFalse(
            ReceitaAluguel.objects.filter(competencia_mes=1).exists()
        )

    def test_periodo_fora_da_vigencia_do_contrato_nao_gera(self):
        """Solicitar um mês fora da vigência do contrato não gera nenhuma receita."""
        criadas, existiam = gerar_receitas_para_contrato(
            self.contrato,
            data_inicio=date(2025, 1, 1),
            data_fim=date(2025, 1, 31),
        )

        self.assertEqual(criadas, 0)
        self.assertEqual(existiam, 0)


class GerarReceitasMesTest(TestCase):
    """Testa a função gerar_receitas_mes (todos os contratos de um mês)."""

    def setUp(self):
        self.imovel, self.locatario = _criar_base()

    def test_geracao_mes_atual(self):
        """gerar_receitas_mes cria receita para contrato ativo no mês atual."""
        hoje = date.today()
        mes = hoje.month
        ano = hoje.year
        ultimo_dia = monthrange(ano, mes)[1]

        contrato = _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(ano, mes, 1),
            data_fim=date(ano, mes, ultimo_dia),
        )

        criadas, existiam = gerar_receitas_mes(mes, ano)

        self.assertEqual(criadas, 1)
        self.assertEqual(existiam, 0)
        self.assertTrue(
            ReceitaAluguel.objects.filter(
                contrato=contrato,
                competencia_mes=mes,
                competencia_ano=ano,
            ).exists()
        )

    def test_contrato_inativo_ignorado_em_gerar_mes(self):
        """Contratos inativos não são incluídos na geração por mês."""
        hoje = date.today()
        _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(hoje.year, hoje.month, 1),
            data_fim=date(hoje.year, hoje.month, monthrange(hoje.year, hoje.month)[1]),
            status='encerrado',
        )

        criadas, _ = gerar_receitas_mes(hoje.month, hoje.year)

        self.assertEqual(criadas, 0)

    def test_multiplos_contratos_no_mes(self):
        """Múltiplos contratos ativos geram uma receita cada."""
        hoje = date.today()
        mes, ano = hoje.month, hoje.year
        ultimo_dia = monthrange(ano, mes)[1]

        imovel2 = Imovel.objects.create(nome='Casa Teste', endereco='Rua B', cidade='SP', estado='SP')

        _criar_contrato(self.imovel, self.locatario, date(ano, mes, 1), date(ano, mes, ultimo_dia))
        _criar_contrato(imovel2, self.locatario, date(ano, mes, 1), date(ano, mes, ultimo_dia))

        criadas, _ = gerar_receitas_mes(mes, ano)

        self.assertEqual(criadas, 2)
        self.assertEqual(ReceitaAluguel.objects.count(), 2)

    def test_sem_contratos_ativos_retorna_zero(self):
        """Sem contratos ativos no período, retorna (0, 0)."""
        criadas, existiam = gerar_receitas_mes(1, 2000)
        self.assertEqual(criadas, 0)
        self.assertEqual(existiam, 0)


class ReceitaAluguelValidacaoTest(TestCase):
    """Testa clean() e save() de ReceitaAluguel."""

    def setUp(self):
        self.imovel, self.locatario = _criar_base()
        self.imovel2 = Imovel.objects.create(nome='Outro Imóvel', endereco='Rua C', cidade='SP', estado='SP')
        self.contrato = _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(2024, 1, 1),
            data_fim=date(2024, 12, 31),
        )

    def test_save_preenche_imovel_automaticamente(self):
        """save() preenche imovel a partir do contrato quando não informado."""
        receita = ReceitaAluguel(
            contrato=self.contrato,
            competencia_mes=1,
            competencia_ano=2024,
            data_vencimento=date(2024, 1, 10),
            valor_previsto=Decimal('2000.00'),
        )
        receita.save()

        self.assertEqual(receita.imovel_id, self.imovel.pk)

    def test_clean_rejeita_imovel_diferente_do_contrato(self):
        """clean() levanta ValidationError quando imovel difere do contrato."""
        from django.core.exceptions import ValidationError

        receita = ReceitaAluguel(
            contrato=self.contrato,
            imovel=self.imovel2,
            competencia_mes=2,
            competencia_ano=2024,
            data_vencimento=date(2024, 2, 10),
            valor_previsto=Decimal('2000.00'),
        )

        with self.assertRaises(ValidationError):
            receita.clean()


class ReceitaSaveForcaImovelTest(TestCase):
    """save() deve sempre forçar imovel a partir do contrato, mesmo se informado outro."""

    def setUp(self):
        self.imovel, self.locatario = _criar_base()
        self.imovel2 = Imovel.objects.create(nome='Outro Imóvel', endereco='Rua C', cidade='SP', estado='SP')
        self.contrato = _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(2024, 1, 1),
            data_fim=date(2024, 12, 31),
        )

    def test_save_sobrescreve_imovel_errado(self):
        """Mesmo passando imovel2, save() deve substituir pelo imóvel do contrato."""
        receita = ReceitaAluguel(
            contrato=self.contrato,
            imovel=self.imovel2,
            competencia_mes=3,
            competencia_ano=2024,
            data_vencimento=date(2024, 3, 10),
            valor_previsto=Decimal('2000.00'),
        )
        receita.save()
        self.assertEqual(receita.imovel_id, self.imovel.pk)


class EstaAtrasadaTest(TestCase):
    """Testa a property esta_atrasada em ReceitaAluguel e Despesa."""

    def setUp(self):
        self.imovel, self.locatario = _criar_base()
        self.contrato = _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(2020, 1, 1),
            data_fim=date(2030, 12, 31),
        )
        self.ontem = timezone.localdate() - timedelta(days=1)
        self.amanha = timezone.localdate() + timedelta(days=1)

    def _receita(self, vencimento, status):
        return ReceitaAluguel(
            contrato=self.contrato,
            imovel=self.imovel,
            competencia_mes=1,
            competencia_ano=2024,
            data_vencimento=vencimento,
            valor_previsto=Decimal('1000.00'),
            status=status,
        )

    def test_receita_prevista_vencida_esta_atrasada(self):
        r = self._receita(self.ontem, 'previsto')
        self.assertTrue(r.esta_atrasada)

    def test_receita_prevista_nao_vencida_nao_esta_atrasada(self):
        r = self._receita(self.amanha, 'previsto')
        self.assertFalse(r.esta_atrasada)

    def test_receita_recebida_nao_esta_atrasada(self):
        r = self._receita(self.ontem, 'recebido')
        r.valor_recebido = Decimal('1000.00')  # quitada de fato — saldo zero
        self.assertFalse(r.esta_atrasada)

    def test_receita_status_recebido_mas_com_saldo_residual_esta_atrasada(self):
        """
        Regra financeira oficial (item 12): não confia isoladamente no
        status. Um status='recebido' sem o valor_recebido correspondente
        (ex.: inconsistência, ou multa lançada depois da baixa sem
        reconsolidação) continua sendo tratado como atrasado, pois há saldo
        em aberto vencido de verdade.
        """
        r = self._receita(self.ontem, 'recebido')  # valor_recebido nunca setado (None)
        self.assertTrue(r.esta_atrasada)
        self.assertEqual(r.saldo_em_aberto, Decimal('1000.00'))

    def test_receita_parcial_vencida_esta_atrasada(self):
        """Parcial NÃO é quitada: vencida com saldo em aberto continua atrasada."""
        r = self._receita(self.ontem, 'parcial')
        r.valor_recebido = Decimal('400.00')
        self.assertTrue(r.esta_atrasada)
        self.assertEqual(r.saldo_em_aberto, Decimal('600.00'))

    def test_receita_parcial_quitada_pelo_valor_nao_esta_atrasada(self):
        r = self._receita(self.ontem, 'parcial')
        r.valor_recebido = Decimal('1000.00')
        self.assertTrue(r.esta_quitada)
        self.assertFalse(r.esta_atrasada)

    def test_receita_cancelada_nao_esta_atrasada(self):
        r = self._receita(self.ontem, 'cancelado')
        self.assertFalse(r.esta_atrasada)

    def test_receita_cancelada_com_valor_residual_nos_campos_nao_esta_atrasada(self):
        """Cancelada sempre tem saldo_em_aberto=0, mesmo com valor_recebido parcial nos campos."""
        r = self._receita(self.ontem, 'cancelado')
        r.valor_recebido = Decimal('400.00')
        self.assertEqual(r.saldo_em_aberto, Decimal('0.00'))
        self.assertFalse(r.esta_atrasada)

    def test_receita_futura_com_saldo_nao_esta_atrasada(self):
        r = self._receita(self.amanha, 'previsto')
        self.assertEqual(r.saldo_em_aberto, Decimal('1000.00'))
        self.assertFalse(r.esta_atrasada)

    def test_despesa_prevista_vencida_esta_atrasada(self):
        d = Despesa(
            imovel=self.imovel,
            descricao='Teste',
            data_vencimento=self.ontem,
            valor=Decimal('500.00'),
            status='prevista',
        )
        self.assertTrue(d.esta_atrasada)

    def test_despesa_paga_nao_esta_atrasada(self):
        d = Despesa(
            imovel=self.imovel,
            descricao='Teste',
            data_vencimento=self.ontem,
            valor=Decimal('500.00'),
            status='paga',
        )
        self.assertFalse(d.esta_atrasada)

    def test_despesa_cancelada_nao_esta_atrasada(self):
        d = Despesa(
            imovel=self.imovel,
            descricao='Teste',
            data_vencimento=self.ontem,
            valor=Decimal('500.00'),
            status='cancelada',
        )
        self.assertFalse(d.esta_atrasada)


class ReceitasInadimplentesQsTest(TestCase):
    """Testa receitas_inadimplentes_qs() — regra de vencimento, não de status."""

    def setUp(self):
        self.imovel, self.locatario = _criar_base()
        self.contrato = _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(2020, 1, 1),
            data_fim=date(2030, 12, 31),
        )
        self.ontem = timezone.localdate() - timedelta(days=1)
        self.amanha = timezone.localdate() + timedelta(days=1)

    def _criar_receita(self, vencimento, status, mes=1, valor_recebido=None):
        return ReceitaAluguel.objects.create(
            contrato=self.contrato,
            imovel=self.imovel,
            competencia_mes=mes,
            competencia_ano=2024,
            data_vencimento=vencimento,
            valor_previsto=Decimal('1000.00'),
            valor_recebido=valor_recebido,
            status=status,
        )

    def test_vencida_previsto_aparece_na_lista(self):
        self._criar_receita(self.ontem, 'previsto')
        self.assertEqual(receitas_inadimplentes_qs().count(), 1)

    def test_vencida_atrasado_aparece_na_lista(self):
        self._criar_receita(self.ontem, 'atrasado')
        self.assertEqual(receitas_inadimplentes_qs().count(), 1)

    def test_vencida_quitada_pelo_saldo_nao_aparece(self):
        # Regra de SALDO: só sai da inadimplência quando o recebido cobre o devido
        self._criar_receita(self.ontem, 'recebido', valor_recebido=Decimal('1000.00'))
        self.assertEqual(receitas_inadimplentes_qs().count(), 0)

    def test_vencida_status_recebido_com_saldo_aparece(self):
        # Status divergente do saldo: a fonte financeira é o saldo → aparece
        self._criar_receita(self.ontem, 'recebido', valor_recebido=Decimal('400.00'))
        self.assertEqual(receitas_inadimplentes_qs().count(), 1)

    def test_nao_vencida_nao_aparece(self):
        self._criar_receita(self.amanha, 'previsto', mes=2)
        self.assertEqual(receitas_inadimplentes_qs().count(), 0)


class ContratoValidacaoTest(TestCase):
    """Testa validações do modelo Contrato."""

    def setUp(self):
        self.imovel, self.locatario = _criar_base()

    def test_dia_vencimento_invalido_falha_na_validacao(self):
        """dia_vencimento = 0 deve levantar ValidationError via full_clean()."""
        contrato = Contrato(
            imovel=self.imovel,
            locatario=self.locatario,
            data_inicio=date(2024, 1, 1),
            data_fim=date(2024, 12, 31),
            valor_aluguel=Decimal('2000.00'),
            dia_vencimento=0,
        )
        with self.assertRaises(ValidationError):
            contrato.full_clean()

    def test_dia_vencimento_maior_que_31_falha(self):
        contrato = Contrato(
            imovel=self.imovel,
            locatario=self.locatario,
            data_inicio=date(2024, 1, 1),
            data_fim=date(2024, 12, 31),
            valor_aluguel=Decimal('2000.00'),
            dia_vencimento=32,
        )
        with self.assertRaises(ValidationError):
            contrato.full_clean()

    def test_contratos_sobrepostos_falham(self):
        """Dois contratos ativos no mesmo imóvel com período sobreposto são bloqueados."""
        _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(2024, 1, 1),
            data_fim=date(2024, 12, 31),
        )
        contrato2 = Contrato(
            imovel=self.imovel,
            locatario=self.locatario,
            data_inicio=date(2024, 6, 1),
            data_fim=date(2025, 5, 31),
            valor_aluguel=Decimal('2500.00'),
            dia_vencimento=10,
            status='ativo',
        )
        with self.assertRaises(ValidationError):
            contrato2.clean()

    def test_contratos_nao_sobrepostos_sao_permitidos(self):
        """Contratos em períodos consecutivos não devem falhar."""
        _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(2024, 1, 1),
            data_fim=date(2024, 6, 30),
        )
        contrato2 = Contrato(
            imovel=self.imovel,
            locatario=self.locatario,
            data_inicio=date(2024, 7, 1),
            data_fim=date(2025, 6, 30),
            valor_aluguel=Decimal('2500.00'),
            dia_vencimento=10,
            status='ativo',
        )
        try:
            contrato2.clean()
        except ValidationError:
            self.fail('clean() não deveria levantar erro para contratos não sobrepostos')


class GerarReceitasViewTest(TestCase):
    """Testa a view gerar_receitas_mes_view — GET não gera, POST gera."""

    def setUp(self):
        self.client = Client()
        self.user = com_leitura(User.objects.create_user('test', password='test'))
        _dar_permissoes(self.user, 'add_receitaaluguel')
        self.client.login(username='test', password='test')
        self.imovel, self.locatario = _criar_base()
        hoje = date.today()
        self.mes = hoje.month
        self.ano = hoje.year
        ultimo_dia = monthrange(self.ano, self.mes)[1]
        self.contrato = _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(self.ano, self.mes, 1),
            data_fim=date(self.ano, self.mes, ultimo_dia),
        )

    def test_get_nao_gera_receitas(self):
        """GET na view de geração apenas exibe o formulário, não cria receitas."""
        response = self.client.get(reverse('gerar_receitas_mes'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(ReceitaAluguel.objects.count(), 0)

    def test_post_gera_receitas_e_redireciona(self):
        """POST na view de geração cria receitas e redireciona para a lista."""
        response = self.client.post(
            reverse('gerar_receitas_mes'),
            data={'mes': self.mes, 'ano': self.ano},
        )
        self.assertRedirects(
            response,
            f"{reverse('receitas_list')}?mes={self.mes}&ano={self.ano}",
        )
        self.assertEqual(ReceitaAluguel.objects.count(), 1)

    def test_get_contagem_inclui_contrato_prazo_indeterminado_apos_data_fim(self):
        """
        A contagem de contratos_ativos exibida na tela GET deve incluir
        contratos por prazo indeterminado mesmo em meses após a data_fim
        original (usa contratos_para_geracao_mes, não data_fim__gte cru).
        """
        contrato_indeterminado = _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(2020, 1, 1), data_fim=date(2020, 12, 31),
            prazo_indeterminado=True,
        )
        mes_futuro, ano_futuro = 6, date.today().year + 1
        response = self.client.get(reverse('gerar_receitas_mes'), {'mes': mes_futuro, 'ano': ano_futuro})
        self.assertGreaterEqual(response.context['contratos_ativos'], 1)

    def test_get_contagem_exclui_contrato_encerrado(self):
        contrato_encerrado = _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(2020, 1, 1), data_fim=date(2030, 12, 31),
            status='encerrado',
        )
        response = self.client.get(reverse('gerar_receitas_mes'), {'mes': self.mes, 'ano': self.ano})
        # Apenas self.contrato (ativo) deve contar; o encerrado nunca conta.
        self.assertEqual(response.context['contratos_ativos'], 1)

    def test_get_contagem_exclui_data_encerramento_real_anterior_ao_mes(self):
        imovel2 = Imovel.objects.create(nome='Imóvel Encerrado', endereco='Rua Z', cidade='SP', estado='SP')
        inicio_do_mes = date(self.ano, self.mes, 1)
        contrato_encerrado_de_fato = _criar_contrato(
            imovel2, self.locatario,
            data_inicio=date(2020, 1, 1), data_fim=date(2030, 12, 31),
            prazo_indeterminado=True,
            data_encerramento_real=inicio_do_mes - timedelta(days=1),
        )
        contratos_aptos = list(contratos_para_geracao_mes(self.mes, self.ano))
        self.assertNotIn(contrato_encerrado_de_fato, contratos_aptos)


class ExportContratosViewTest(TestCase):
    """Testa filtros de status e imóvel na exportação de contratos."""

    def setUp(self):
        self.client = Client()
        self.user = com_leitura(User.objects.create_user('test2', password='test2'))
        self.client.login(username='test2', password='test2')
        self.imovel, self.locatario = _criar_base()
        self.imovel2 = Imovel.objects.create(nome='Casa Dois', endereco='Rua D', cidade='SP', estado='SP')
        _criar_contrato(self.imovel, self.locatario, date(2024, 1, 1), date(2024, 12, 31), status='ativo')
        _criar_contrato(self.imovel2, self.locatario, date(2023, 1, 1), date(2023, 12, 31), status='encerrado')

    def test_sem_filtro_exporta_apenas_ativos(self):
        response = self.client.get(reverse('export_contratos', args=['csv']))
        self.assertEqual(response.status_code, 200)
        content = response.content.decode('utf-8-sig')
        self.assertIn('Apartamento Teste', content)
        self.assertNotIn('Casa Dois', content)

    def test_filtro_encerrado_exporta_encerrados(self):
        response = self.client.get(reverse('export_contratos', args=['csv']) + '?status=encerrado')
        self.assertEqual(response.status_code, 200)
        content = response.content.decode('utf-8-sig')
        self.assertIn('Casa Dois', content)
        self.assertNotIn('Apartamento Teste', content)

    def test_filtro_imovel_restringe_resultado(self):
        response = self.client.get(
            reverse('export_contratos', args=['csv']) + f'?status=ativo&imovel={self.imovel.pk}'
        )
        self.assertEqual(response.status_code, 200)
        content = response.content.decode('utf-8-sig')
        self.assertIn('Apartamento Teste', content)
        self.assertNotIn('Casa Dois', content)


# ─── Testes da tela de baixa de aluguéis ──────────────────────────────────────

class BaixaReceitasViewTest(TestCase):
    """Testa a view baixa_receitas_mes_view."""

    def setUp(self):
        self.client = Client()
        self.user = com_leitura(User.objects.create_user('baixa_user', password='pass'))
        _dar_permissoes(self.user, 'add_receitaaluguel', 'change_receitaaluguel')
        self.imovel, self.locatario = _criar_base()
        hoje = date.today()
        self.mes = hoje.month
        self.ano = hoje.year
        ultimo_dia = monthrange(self.ano, self.mes)[1]
        self.contrato = _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(self.ano, self.mes, 1),
            data_fim=date(self.ano, self.mes, ultimo_dia),
        )
        self.receita = ReceitaAluguel.objects.create(
            contrato=self.contrato,
            imovel=self.imovel,
            competencia_mes=self.mes,
            competencia_ano=self.ano,
            data_vencimento=date(self.ano, self.mes, 10),
            valor_previsto=Decimal('2000.00'),
            status='previsto',
        )

    def test_exige_login(self):
        """Acesso sem login deve redirecionar para login."""
        response = self.client.get(reverse('baixa_receitas_mes'))
        self.assertEqual(response.status_code, 302)
        self.assertIn('/login/', response['Location'])

    def test_get_mostra_receitas_do_mes(self):
        """GET autenticado mostra as receitas do mês selecionado."""
        self.client.login(username='baixa_user', password='pass')
        response = self.client.get(
            reverse('baixa_receitas_mes'),
            {'mes': self.mes, 'ano': self.ano},
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(self.receita, response.context['receitas'])

    def test_post_marcar_recebida_preenche_valores(self):
        """POST action=marcar_recebida preenche valor e data automaticamente."""
        self.client.login(username='baixa_user', password='pass')
        self.client.post(
            reverse('baixa_receitas_mes'),
            {
                'mes': self.mes, 'ano': self.ano,
                'receita_id': self.receita.pk,
                'action': 'marcar_recebida',
            },
        )
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.status, 'recebido')
        self.assertEqual(self.receita.valor_recebido, Decimal('2000.00'))
        self.assertEqual(self.receita.data_recebimento, date.today())

    def test_post_marcar_recebida_completa_o_saldo(self):
        """
        marcar_recebida agora registra um RecebimentoReceita do SALDO restante —
        o valor parcial anterior (1800) é preservado como recebimento histórico
        e o total consolidado passa a cobrir o previsto (2000).
        """
        self.receita.valor_recebido = Decimal('1800.00')
        self.receita.save()
        self.client.login(username='baixa_user', password='pass')
        self.client.post(
            reverse('baixa_receitas_mes'),
            {
                'mes': self.mes, 'ano': self.ano,
                'receita_id': self.receita.pk,
                'action': 'marcar_recebida',
            },
        )
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.valor_recebido, Decimal('2000.00'))
        self.assertEqual(self.receita.status, 'recebido')
        # dois recebimentos: o legado materializado (1800) + o do saldo (200)
        self.assertEqual(self.receita.recebimentos.count(), 2)
        valores = sorted(self.receita.recebimentos.values_list('valor', flat=True))
        self.assertEqual(valores, [Decimal('200.00'), Decimal('1800.00')])

    def test_post_registrar_recebimento_parcial(self):
        """
        POST action=registrar_recebimento cria um RecebimentoReceita e a
        consolidação (valor_recebido/status) é derivada dele — não há mais
        edição direta dos campos consolidados na tela de baixa.
        """
        self.client.login(username='baixa_user', password='pass')
        self.client.post(
            reverse('baixa_receitas_mes'),
            {
                'mes': self.mes, 'ano': self.ano,
                'receita_id': self.receita.pk,
                'action': 'registrar_recebimento',
                'valor': '1500.00',
                'data_recebimento': date.today().isoformat(),
                'observacoes': 'Pagamento parcial acordado.',
            },
        )
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.status, 'parcial')
        self.assertEqual(self.receita.valor_recebido, Decimal('1500.00'))
        recebimento = self.receita.recebimentos.get()
        self.assertEqual(recebimento.valor, Decimal('1500.00'))
        self.assertEqual(recebimento.origem, 'manual')
        self.assertEqual(recebimento.observacoes, 'Pagamento parcial acordado.')
        self.assertEqual(recebimento.criado_por.username, 'baixa_user')


# ─── Testes do relatório de contabilidade ────────────────────────────────────

class RelatorioContabilidadeTest(TestCase):
    """Testa exportar_relatorio_contabilidade_xlsx."""

    def setUp(self):
        self.client = Client()
        self.user = com_leitura(User.objects.create_user('cont_user', password='pass'))
        self.client.login(username='cont_user', password='pass')
        self.imovel, self.locatario = _criar_base()
        self.contrato = _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(2024, 1, 1),
            data_fim=date(2024, 12, 31),
        )
        self.ontem = timezone.localdate() - timedelta(days=1)

    def test_retorna_xlsx_valido(self):
        """GET retorna status 200 e content-type XLSX."""
        response = self.client.get(
            reverse('export_relatorio_contabilidade'),
            {'mes': 1, 'ano': 2024},
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml', response['Content-Type'])

    def test_inclui_receita_vencida_com_status_previsto(self):
        """Receita vencida mas com status=previsto aparece na aba Inadimplência."""
        import io
        try:
            import openpyxl
        except ImportError:
            self.skipTest('openpyxl não instalado')

        ReceitaAluguel.objects.create(
            contrato=self.contrato,
            imovel=self.imovel,
            competencia_mes=1,
            competencia_ano=2024,
            data_vencimento=self.ontem,
            valor_previsto=Decimal('2000.00'),
            status='previsto',
        )
        response = self.client.get(
            reverse('export_relatorio_contabilidade'),
            {'mes': 1, 'ano': 2024},
        )
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertIn('Inadimplência Aberta', wb.sheetnames)
        ws = wb['Inadimplência Aberta']
        valores = [row[0] for row in ws.iter_rows(min_row=2, values_only=True) if row[0]]
        self.assertTrue(any(self.imovel.nome in str(v) for v in valores))


class RelatoriosExportacoesPermissoesTest(TestCase):
    """Item 4 (rodada pós-revisão): permissões por tipo de exportação."""

    def setUp(self):
        from django.contrib.auth.models import Permission
        self.Permission = Permission
        self.client = Client()
        self.imovel, self.locatario = _criar_base()
        self.contrato = _criar_contrato(
            self.imovel, self.locatario, data_inicio=date(2024, 1, 1), data_fim=date(2024, 12, 31),
        )
        Despesa.objects.create(
            imovel=self.imovel, categoria='iptu', descricao='IPTU', data_vencimento=date(2024, 3, 10),
            data_pagamento=date(2024, 3, 10), valor=Decimal('200.00'), status='paga',
        )

    def _usuario(self, nome, *perms):
        user = User.objects.create_user(nome, password='pass')
        if perms:
            user.user_permissions.add(*self.Permission.objects.filter(codename__in=perms))
        self.client.login(username=nome, password='pass')
        return user

    def test_usuario_so_de_receitas_nao_exporta_despesas(self):
        self._usuario('rel_so_rec_1', 'view_receitaaluguel')
        response = self.client.get(reverse('export_despesas', args=['csv']), {'mes': 3, 'ano': 2024})
        self.assertEqual(response.status_code, 403)

    def test_usuario_so_de_receitas_nao_exporta_relatorio_mensal_completo(self):
        self._usuario('rel_so_rec_2', 'view_receitaaluguel')
        response = self.client.get(reverse('export_relatorio_mensal', args=['xlsx']), {'mes': 3, 'ano': 2024})
        self.assertEqual(response.status_code, 403)

    def test_usuario_sem_documentos_nao_exporta_relatorio_contabil_completo(self):
        self._usuario('rel_sem_docs', 'view_receitaaluguel', 'view_despesa')
        response = self.client.get(reverse('export_relatorio_contabilidade'), {'mes': 3, 'ano': 2024})
        self.assertEqual(response.status_code, 403)

    def test_usuario_com_todas_as_permissoes_exporta_todas_as_abas(self):
        self._usuario(
            'rel_completo', 'view_receitaaluguel', 'view_despesa', 'view_documento',
        )
        response = self.client.get(reverse('export_relatorio_mensal', args=['xlsx']), {'mes': 3, 'ano': 2024})
        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml', response['Content-Type'])
        response = self.client.get(reverse('export_relatorio_contabilidade'), {'mes': 3, 'ano': 2024})
        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml', response['Content-Type'])

    def test_fornecedor_nao_aparece_para_usuario_sem_view_despesa(self):
        """Sem view_despesa, o relatório mensal completo (que tem a aba Despesas/Fornecedor) é inacessível."""
        self._usuario('rel_sem_desp', 'view_receitaaluguel')
        response = self.client.get(reverse('export_relatorio_mensal', args=['xlsx']), {'mes': 3, 'ano': 2024})
        self.assertEqual(response.status_code, 403)

    def test_documentos_nao_aparecem_para_usuario_sem_view_documento(self):
        """Sem view_documento, o relatório contábil (aba Docs. Pendentes) é inacessível."""
        self._usuario('rel_sem_doc_view', 'view_receitaaluguel', 'view_despesa')
        response = self.client.get(reverse('export_relatorio_contabilidade'), {'mes': 3, 'ano': 2024})
        self.assertEqual(response.status_code, 403)

    def test_tela_relatorios_esconde_cards_sem_permissao(self):
        self._usuario('rel_tela_so_rec', 'view_receitaaluguel')
        response = self.client.get(reverse('relatorios'))
        self.assertEqual(response.status_code, 200)
        conteudo = response.content.decode()
        self.assertIn('Receitas por Período', conteudo)
        self.assertNotIn('Despesas por Período', conteudo)
        self.assertNotIn('Relatório para Contabilidade', conteudo)
        self.assertIn('exige acesso a receitas', conteudo)

    def test_tela_relatorios_sem_nenhuma_permissao_403(self):
        self._usuario('rel_tela_sem_perm')
        response = self.client.get(reverse('relatorios'))
        self.assertEqual(response.status_code, 403)


# ─── Testes de DocumentoObrigatorio ──────────────────────────────────────────

class DocumentoObrigatorioTest(TestCase):
    """Testa o model DocumentoObrigatorio e contadores."""

    def setUp(self):
        self.client = Client()
        self.user = com_leitura(User.objects.create_user('doc_user', password='pass'))
        self.client.login(username='doc_user', password='pass')
        self.imovel, self.locatario = _criar_base()

    def test_pendente_sem_documento_vinculado(self):
        """DocumentoObrigatorio sem documento vinculado deve ser pendente."""
        from documentos.models import DocumentoObrigatorio
        ob = DocumentoObrigatorio(
            imovel=self.imovel,
            tipo='matricula',
            descricao='Matrícula do imóvel',
            obrigatorio=True,
        )
        self.assertTrue(ob.pendente)

    def test_nao_pendente_com_documento_vinculado(self):
        """DocumentoObrigatorio com documento vinculado não é pendente."""
        from documentos.models import DocumentoObrigatorio, Documento
        import tempfile, os
        from django.core.files.uploadedfile import SimpleUploadedFile
        doc = Documento.objects.create(
            titulo='Matrícula',
            tipo='matricula',
            arquivo=SimpleUploadedFile('test.pdf', b'conteudo'),
            imovel=self.imovel,
        )
        ob = DocumentoObrigatorio(
            imovel=self.imovel,
            tipo='matricula',
            descricao='Matrícula do imóvel',
            obrigatorio=True,
            documento=doc,
        )
        self.assertFalse(ob.pendente)

    def test_dashboard_conta_pendentes(self):
        """Dashboard exibe contador de docs obrigatórios pendentes."""
        from documentos.models import DocumentoObrigatorio
        DocumentoObrigatorio.objects.create(
            imovel=self.imovel,
            tipo='matricula',
            descricao='Matrícula',
            obrigatorio=True,
        )
        response = self.client.get(reverse('dashboard'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['docs_obrigatorios_pendentes'], 1)

    def test_dashboard_nao_obrigatorio_nao_conta(self):
        """DocumentoObrigatorio com obrigatorio=False não conta no dashboard."""
        from documentos.models import DocumentoObrigatorio
        DocumentoObrigatorio.objects.create(
            imovel=self.imovel,
            tipo='seguro',
            descricao='Seguro',
            obrigatorio=False,
        )
        response = self.client.get(reverse('dashboard'))
        self.assertEqual(response.context['docs_obrigatorios_pendentes'], 0)


# ─── Teste do backup_local (ZIP) ──────────────────────────────────────────────

class BackupLocalTest(TestCase):
    """Testa o comando backup_local — geração de ZIP."""

    def test_gera_arquivo_zip_com_manifest(self):
        """backup_local deve criar um .zip contendo manifest.txt."""
        import tempfile
        import zipfile as ziplib
        from io import StringIO
        from django.core.management import call_command

        with tempfile.TemporaryDirectory() as tmpdir:
            call_command('backup_local', '--destino', tmpdir, verbosity=0)
            zips = [f for f in os.listdir(tmpdir) if f.endswith('.zip')]
            self.assertEqual(len(zips), 1, 'Deve existir exatamente 1 arquivo .zip')
            with ziplib.ZipFile(os.path.join(tmpdir, zips[0])) as zf:
                self.assertIn('manifest.txt', zf.namelist())

    def test_manter_apenas_n_backups(self):
        """--manter N apaga os zips mais antigos mantendo apenas N."""
        import tempfile
        import zipfile as ziplib
        from django.core.management import call_command

        with tempfile.TemporaryDirectory() as tmpdir:
            # Cria 3 backups antigos com timestamps distintos
            for i in range(1, 4):
                path = os.path.join(tmpdir, f'backup_2024010{i}_120000.zip')
                with ziplib.ZipFile(path, 'w') as zf:
                    zf.writestr('manifest.txt', f'dummy {i}')
            # Executa com --manter 2: cria +1, remove os 2 mais antigos
            call_command('backup_local', '--destino', tmpdir, '--manter', '2', verbosity=0)
            zips = [f for f in os.listdir(tmpdir) if f.endswith('.zip')]
            self.assertEqual(len(zips), 2)


# ─── Testes do formulário baixa_receitas (Items 1-3) ─────────────────────────

class BaixaReceitasFormTest(TestCase):
    """Verifica que baixa_receitas.html não tem formulário aninhado."""

    def setUp(self):
        self.client = Client()
        user = com_leitura(User.objects.create_user('form_user', password='pass'))
        _dar_permissoes(user, 'add_receitaaluguel', 'change_receitaaluguel')
        self.client.login(username='form_user', password='pass')

    def test_get_renderiza_sem_form_aninhado(self):
        """GET da tela de baixa deve retornar status 200 sem form aninhado."""
        response = self.client.get(reverse('baixa_receitas_mes'))
        self.assertEqual(response.status_code, 200)
        content = response.content.decode()
        # Não deve haver dois <form method="post"> dentro de um <form method="get">
        # Verificamos que o HTML contém o botão Gerar mas como form separado
        self.assertIn('gerar_receitas', content)
        self.assertIn('Gerar Receitas Esperadas', content)

    def test_atributos_decimais_nao_localizados_em_pt_br(self):
        """
        Item 11: com LANGUAGE_CODE='pt-br' (vírgula decimal na exibição), os
        atributos data-* usados pelo JS (parseFloat/valor de <input>) devem
        continuar com ponto — só o texto para o usuário usa vírgula.
        """
        imovel, locatario = _criar_base()
        contrato = _criar_contrato(
            imovel, locatario, date(2024, 1, 1), date(2024, 12, 31),
            valor_aluguel=Decimal('1234.56'),
        )
        receita = ReceitaAluguel.objects.create(
            contrato=contrato, imovel=imovel, competencia_mes=3, competencia_ano=2024,
            data_vencimento=date(2024, 3, 10), valor_previsto=Decimal('1234.56'),
            multa=Decimal('12.34'), juros=Decimal('5.67'), status='previsto',
        )
        response = self.client.get(reverse('baixa_receitas_mes'), {'mes': 3, 'ano': 2024})
        content = response.content.decode()
        # saldo_em_aberto = previsto + multa + juros - desconto = 1234.56 + 12.34 + 5.67
        self.assertIn('data-saldo="1252.57"', content)
        self.assertIn('data-multa="12.34"', content)
        self.assertIn('data-juros="5.67"', content)
        # nunca com vírgula nesses atributos (isso quebraria o JS)
        self.assertNotIn('data-saldo="1252,57"', content)
        self.assertNotIn('data-multa="12,34"', content)
        self.assertNotIn('data-juros="5,67"', content)

    def test_contadores_distinguem_quitadas_parciais_canceladas(self):
        """Item 13: o contador da tela de baixa não deve lumpar 'parcial' em 'Recebidas'."""
        imovel, locatario = _criar_base()
        contrato = _criar_contrato(imovel, locatario, date(2024, 1, 1), date(2024, 12, 31))
        ReceitaAluguel.objects.create(
            contrato=contrato, imovel=imovel, competencia_mes=3, competencia_ano=2024,
            data_vencimento=date(2024, 3, 10), valor_previsto=Decimal('1000.00'), status='previsto',
        )
        response = self.client.get(reverse('baixa_receitas_mes'), {'mes': 3, 'ano': 2024})
        content = response.content.decode()
        self.assertIn('id="cnt-quitadas"', content)
        self.assertIn('id="cnt-parciais"', content)
        self.assertIn('id="cnt-canceladas"', content)
        self.assertIn('Quitadas', content)
        self.assertIn('Parciais', content)
        self.assertIn('Canceladas', content)
        # JS não deve mais contar 'parcial' como recebida
        self.assertNotIn("s === 'recebido' || s === 'parcial'", content)

    def test_post_gerar_receitas_a_partir_de_baixa(self):
        """POST action=gerar_receitas na tela de baixa gera receitas corretamente."""
        imovel, locatario = _criar_base()
        hoje = date.today()
        ultimo_dia = monthrange(hoje.year, hoje.month)[1]
        _criar_contrato(
            imovel, locatario,
            date(hoje.year, hoje.month, 1),
            date(hoje.year, hoje.month, ultimo_dia),
        )
        response = self.client.post(
            reverse('baixa_receitas_mes'),
            {'mes': hoje.month, 'ano': hoje.year, 'action': 'gerar_receitas'},
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(ReceitaAluguel.objects.count(), 1)


# ─── Testes de DocumentoObrigatorio.clean() (Item 4) ─────────────────────────

class DocumentoObrigatorioCleanTest(TestCase):
    """Testa o método clean() do DocumentoObrigatorio."""

    def setUp(self):
        from documentos.models import Documento, DocumentoObrigatorio
        from django.core.files.uploadedfile import SimpleUploadedFile
        self.imovel, _ = _criar_base()
        self.imovel2 = Imovel.objects.create(nome='Outro', endereco='Rua Z', cidade='SP', estado='SP')
        self.doc_imovel1 = Documento.objects.create(
            titulo='Matrícula 1',
            tipo='matricula',
            arquivo=SimpleUploadedFile('m1.pdf', b'x'),
            imovel=self.imovel,
        )
        self.doc_sem_imovel = Documento.objects.create(
            titulo='Sem Imóvel',
            tipo='outro',
            arquivo=SimpleUploadedFile('s.pdf', b'x'),
            imovel=None,
        )

    def test_clean_aceita_documento_do_mesmo_imovel(self):
        """clean() não levanta erro quando documento pertence ao mesmo imóvel."""
        from documentos.models import DocumentoObrigatorio
        from django.core.exceptions import ValidationError
        ob = DocumentoObrigatorio(
            imovel=self.imovel,
            tipo='matricula',
            descricao='Matrícula',
            documento=self.doc_imovel1,
        )
        try:
            ob.clean()
        except ValidationError:
            self.fail('clean() não deveria levantar ValidationError para documento do mesmo imóvel.')

    def test_clean_rejeita_documento_de_outro_imovel(self):
        """clean() rejeita documento que pertence a outro imóvel."""
        from documentos.models import DocumentoObrigatorio
        from django.core.exceptions import ValidationError
        ob = DocumentoObrigatorio(
            imovel=self.imovel2,
            tipo='matricula',
            descricao='Matrícula',
            documento=self.doc_imovel1,
        )
        with self.assertRaises(ValidationError):
            ob.clean()

    def test_clean_rejeita_documento_sem_imovel(self):
        """clean() rejeita documento sem imóvel vinculado."""
        from documentos.models import DocumentoObrigatorio
        from django.core.exceptions import ValidationError
        ob = DocumentoObrigatorio(
            imovel=self.imovel,
            tipo='outro',
            descricao='Genérico',
            documento=self.doc_sem_imovel,
        )
        with self.assertRaises(ValidationError):
            ob.clean()

    def test_clean_sem_documento_passa(self):
        """clean() sem documento vinculado não levanta erro."""
        from documentos.models import DocumentoObrigatorio
        ob = DocumentoObrigatorio(
            imovel=self.imovel,
            tipo='seguro',
            descricao='Seguro',
            documento=None,
        )
        try:
            ob.clean()
        except Exception:
            self.fail('clean() sem documento não deve levantar exceção.')


# ─── Testes da constraint única (Item 5) ─────────────────────────────────────

class DocumentoObrigatorioUniqueTest(TestCase):
    """Testa que unique_together inclui descricao (imovel, tipo, descricao)."""

    def setUp(self):
        from documentos.models import DocumentoObrigatorio
        self.imovel, _ = _criar_base()
        DocumentoObrigatorio.objects.create(
            imovel=self.imovel,
            tipo='matricula',
            descricao='Matrícula principal',
        )

    def test_mesmo_tipo_descricao_diferente_permitido(self):
        """Dois DocumentoObrigatorio com mesmo tipo mas descrição diferente são permitidos."""
        from documentos.models import DocumentoObrigatorio
        from django.db import IntegrityError
        try:
            DocumentoObrigatorio.objects.create(
                imovel=self.imovel,
                tipo='matricula',
                descricao='Matrícula atualizada 2024',
            )
        except IntegrityError:
            self.fail('Deve ser possível ter mesmo tipo com descrição diferente.')

    def test_mesmo_tipo_mesma_descricao_falha(self):
        """Duplicar imovel+tipo+descricao deve levantar IntegrityError."""
        from documentos.models import DocumentoObrigatorio
        from django.db import IntegrityError
        with self.assertRaises(IntegrityError):
            DocumentoObrigatorio.objects.create(
                imovel=self.imovel,
                tipo='matricula',
                descricao='Matrícula principal',
            )


# ─── Testes do relatório contabilidade — nota inadimplência (Item 6) ──────────

class RelatorioContabilidadeNotaTest(TestCase):
    """Testa que o Resumo do relatório contabilidade inclui nota sobre inadimplência."""

    def setUp(self):
        self.client = Client()
        com_leitura(User.objects.create_user('nota_user', password='pass'))
        self.client.login(username='nota_user', password='pass')

    def test_resumo_contem_nota_inadimplencia(self):
        """A aba Resumo deve conter a nota sobre o escopo da aba Inadimplência Aberta."""
        import io
        try:
            import openpyxl
        except ImportError:
            self.skipTest('openpyxl não instalado')
        response = self.client.get(
            reverse('export_relatorio_contabilidade'),
            {'mes': 1, 'ano': 2024},
        )
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb['Resumo']
        valores = [str(row[0] or '') for row in ws.iter_rows(values_only=True)]
        self.assertTrue(any('Nota' in v for v in valores), 'Resumo deve conter linha de nota.')

    def test_aba_inadimplencia_renomeada(self):
        """A aba de inadimplência deve se chamar 'Inadimplência Aberta'."""
        import io
        try:
            import openpyxl
        except ImportError:
            self.skipTest('openpyxl não instalado')
        response = self.client.get(
            reverse('export_relatorio_contabilidade'),
            {'mes': 1, 'ano': 2024},
        )
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertIn('Inadimplência Aberta', wb.sheetnames)
        self.assertNotIn('Inadimplência', wb.sheetnames)


# ─── Testes do checklist mensal (Item 8) ─────────────────────────────────────

class ChecklistMensalViewTest(TestCase):
    """Testa a view checklist_mensal_view."""

    def setUp(self):
        self.client = Client()
        self.user = com_leitura(User.objects.create_user('chk_user', password='pass'))
        _dar_permissoes(self.user, 'change_fechamentomensal')
        self.imovel, self.locatario = _criar_base()
        hoje = date.today()
        self.mes = hoje.month
        self.ano = hoje.year

    def test_exige_login(self):
        """Acesso sem login deve redirecionar."""
        response = self.client.get(reverse('checklist_mensal'))
        self.assertEqual(response.status_code, 302)
        self.assertIn('/login/', response['Location'])

    def test_get_renderiza_checklist(self):
        """GET autenticado renderiza a página de checklist."""
        self.client.login(username='chk_user', password='pass')
        response = self.client.get(
            reverse('checklist_mensal'),
            {'mes': self.mes, 'ano': self.ano},
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn('checklist', response.context)
        self.assertEqual(len(response.context['checklist']), 11)

    def test_post_marcar_enviado_cria_fechamento(self):
        """POST action=marcar_enviado cria FechamentoMensal e redireciona."""
        from financeiro.models import FechamentoMensal
        self.client.login(username='chk_user', password='pass')
        response = self.client.post(
            reverse('checklist_mensal'),
            {'mes': self.mes, 'ano': self.ano, 'action': 'marcar_enviado'},
        )
        self.assertEqual(response.status_code, 302)
        fechamento = FechamentoMensal.objects.get(mes=self.mes, ano=self.ano)
        self.assertTrue(fechamento.enviado_contabilidade)
        self.assertIsNotNone(fechamento.data_envio_contabilidade)

    def test_post_marcar_enviado_idempotente(self):
        """Chamar marcar_enviado duas vezes não duplica FechamentoMensal."""
        from financeiro.models import FechamentoMensal
        self.client.login(username='chk_user', password='pass')
        for _ in range(2):
            self.client.post(
                reverse('checklist_mensal'),
                {'mes': self.mes, 'ano': self.ano, 'action': 'marcar_enviado'},
            )
        self.assertEqual(FechamentoMensal.objects.filter(mes=self.mes, ano=self.ano).count(), 1)

    def test_dashboard_contem_link_checklist(self):
        """Dashboard deve ter link para o checklist mensal."""
        self.client.login(username='chk_user', password='pass')
        response = self.client.get(reverse('dashboard'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, reverse('checklist_mensal'))

    def test_checklist_sem_receitas_etapa_nao_ok(self):
        """Com nenhuma receita no mês, a etapa 'Receitas geradas' deve estar não-ok."""
        self.client.login(username='chk_user', password='pass')
        response = self.client.get(
            reverse('checklist_mensal'),
            {'mes': self.mes, 'ano': self.ano},
        )
        checklist = response.context['checklist']
        etapa_receitas = next(e for e in checklist if e['item'] == 'Receitas geradas')
        self.assertFalse(etapa_receitas['ok'])

    def test_checklist_receita_cancelada_sem_pendencia_nao_mostra_0_de_1(self):
        """
        Item 13: uma única receita CANCELADA sem saldo em aberto não deve
        gerar a mensagem confusa "0/1 recebida(s)" (soa como pendência que
        não existe) — o detalhe discrimina recebidas/canceladas/em aberto.
        """
        contrato = _criar_contrato(
            self.imovel, self.locatario, date(2020, 1, 1), date(2030, 12, 31),
        )
        receita = ReceitaAluguel.objects.create(
            contrato=contrato, imovel=self.imovel,
            competencia_mes=self.mes, competencia_ano=self.ano,
            data_vencimento=date(self.ano, self.mes, 10), valor_previsto=Decimal('1000.00'),
            status='previsto',
        )
        receita.cancelar()

        self.client.login(username='chk_user', password='pass')
        response = self.client.get(reverse('checklist_mensal'), {'mes': self.mes, 'ano': self.ano})
        checklist = response.context['checklist']
        etapa = next(e for e in checklist if e['item'] == 'Recebimentos confirmados')
        self.assertTrue(etapa['ok'])
        self.assertNotIn('0/1', etapa['detalhe'])
        self.assertIn('1 cancelada(s)', etapa['detalhe'])
        self.assertIn('0 em aberto', etapa['detalhe'])
        self.assertEqual(response.context['receitas_canceladas'], 1)
        self.assertEqual(response.context['receitas_pendentes'], 0)

    def test_checklist_contem_etapa_docs_obrigatorios(self):
        """Checklist deve incluir etapa de documentos obrigatórios."""
        self.client.login(username='chk_user', password='pass')
        response = self.client.get(reverse('checklist_mensal'), {'mes': self.mes, 'ano': self.ano})
        checklist = response.context['checklist']
        items = [e['item'] for e in checklist]
        self.assertIn('Documentos obrigatórios revisados', items)

    def test_docs_obrigatorios_pendentes_torna_etapa_nao_ok(self):
        """Criar DocumentoObrigatorio sem doc vinculado torna a etapa não-ok."""
        from documentos.models import DocumentoObrigatorio
        DocumentoObrigatorio.objects.create(
            imovel=self.imovel, tipo='matricula', descricao='Matrícula', obrigatorio=True
        )
        self.client.login(username='chk_user', password='pass')
        response = self.client.get(reverse('checklist_mensal'), {'mes': self.mes, 'ano': self.ano})
        checklist = response.context['checklist']
        etapa = next(e for e in checklist if e['item'] == 'Documentos obrigatórios revisados')
        self.assertFalse(etapa['ok'])
        self.assertEqual(response.context['docs_obrigatorios_pendentes'], 1)

    def test_docs_obrigatorios_todos_vinculados_torna_etapa_ok(self):
        """DocumentoObrigatorio com doc vinculado torna a etapa ok."""
        from documentos.models import DocumentoObrigatorio, Documento
        from django.core.files.uploadedfile import SimpleUploadedFile
        doc = Documento.objects.create(
            titulo='Matrícula',
            tipo='matricula',
            arquivo=SimpleUploadedFile('m.pdf', b'x'),
            imovel=self.imovel,
        )
        DocumentoObrigatorio.objects.create(
            imovel=self.imovel, tipo='matricula', descricao='Matrícula', obrigatorio=True, documento=doc
        )
        self.client.login(username='chk_user', password='pass')
        response = self.client.get(reverse('checklist_mensal'), {'mes': self.mes, 'ano': self.ano})
        checklist = response.context['checklist']
        etapa = next(e for e in checklist if e['item'] == 'Documentos obrigatórios revisados')
        self.assertTrue(etapa['ok'])

    def test_checklist_contem_etapa_relatorio_contabil(self):
        """Checklist deve incluir etapa informativa de relatório contábil."""
        self.client.login(username='chk_user', password='pass')
        response = self.client.get(reverse('checklist_mensal'), {'mes': self.mes, 'ano': self.ano})
        checklist = response.context['checklist']
        etapa = next((e for e in checklist if e['item'] == 'Relatório contábil exportado'), None)
        self.assertIsNotNone(etapa)
        self.assertTrue(etapa.get('informativa', False))
        self.assertIn(reverse('export_relatorio_contabilidade'), etapa['link'])


class ChecklistPermissoesTest(TestCase):
    """Item 5 (rodada pós-revisão): cada etapa do checklist respeita sua permissão."""

    def setUp(self):
        from django.contrib.auth.models import Permission
        self.Permission = Permission
        self.client = Client()
        hoje = timezone.localdate()
        self.mes, self.ano = hoje.month, hoje.year
        self.imovel, self.locatario = _criar_base()
        self.contrato = _criar_contrato(
            self.imovel, self.locatario, data_inicio=date(2020, 1, 1), data_fim=date(2030, 12, 31),
        )

    def _usuario(self, nome, *perms):
        user = User.objects.create_user(nome, password='pass')
        if perms:
            user.user_permissions.add(*self.Permission.objects.filter(codename__in=perms))
        self.client.login(username=nome, password='pass')
        return user

    def _get(self):
        return self.client.get(reverse('checklist_mensal'), {'mes': self.mes, 'ano': self.ano})

    def test_sem_nenhuma_permissao_403(self):
        self._usuario('chk_sem_perm')
        self.assertEqual(self._get().status_code, 403)

    def test_apenas_receitas_mostra_so_etapas_de_receita(self):
        self._usuario('chk_so_receitas', 'view_receitaaluguel')
        resposta = self._get()
        self.assertEqual(resposta.status_code, 200)
        itens = [e['item'] for e in resposta.context['checklist']]
        self.assertIn('Receitas geradas', itens)
        self.assertIn('Recebimentos confirmados', itens)
        self.assertIn('Inadimplência em dia', itens)
        self.assertNotIn('Despesas pagas', itens)
        self.assertNotIn('Documentos enviados à contabilidade', itens)
        self.assertNotIn('Documentos obrigatórios revisados', itens)
        self.assertNotIn('Reajustes em dia', itens)
        self.assertNotIn('Extrato bancário conciliado', itens)
        # relatório contábil completo exige despesas+documentos também — não aparece
        self.assertNotIn('Relatório contábil exportado', itens)
        self.assertNotIn('Fechamento registrado e enviado', itens)
        conteudo = resposta.content.decode()
        self.assertNotIn('Despesas</div>', conteudo)
        # não deve haver link para a tela de despesas (403 se clicado)
        self.assertNotIn(reverse('despesas_list'), conteudo)

    def test_apenas_despesas_mostra_so_etapa_de_despesa(self):
        self._usuario('chk_so_despesas', 'view_despesa')
        resposta = self._get()
        self.assertEqual(resposta.status_code, 200)
        itens = [e['item'] for e in resposta.context['checklist']]
        self.assertEqual(itens, ['Despesas pagas'])
        self.assertIsNone(resposta.context['total_receitas'])
        self.assertIsNone(resposta.context['inadimplentes'])

    def test_apenas_documentos_mostra_etapas_de_documento(self):
        self._usuario('chk_so_docs', 'view_documento')
        resposta = self._get()
        self.assertEqual(resposta.status_code, 200)
        itens = [e['item'] for e in resposta.context['checklist']]
        self.assertIn('Documentos enviados à contabilidade', itens)
        self.assertIn('Validade dos documentos em dia', itens)
        self.assertNotIn('Documentos obrigatórios revisados', itens)
        self.assertNotIn('Receitas geradas', itens)

    def test_apenas_doc_obrigatorios_mostra_essa_etapa_com_link_so_para_staff(self):
        user = self._usuario('chk_so_doc_obrig', 'view_documentoobrigatorio')
        resposta = self._get()
        itens = [e['item'] for e in resposta.context['checklist']]
        self.assertEqual(itens, ['Documentos obrigatórios revisados'])
        etapa = resposta.context['checklist'][0]
        self.assertIsNone(etapa['link'])  # usuário não é staff — sem link para o Admin
        user.is_staff = True
        user.save()
        resposta2 = self._get()
        etapa2 = resposta2.context['checklist'][0]
        self.assertIsNotNone(etapa2['link'])

    def test_apenas_contratos_mostra_etapa_de_reajustes(self):
        self._usuario('chk_so_contratos', 'view_contrato')
        resposta = self._get()
        itens = [e['item'] for e in resposta.context['checklist']]
        self.assertEqual(itens, ['Reajustes em dia'])

    def test_apenas_conciliacao_mostra_etapa_de_extrato(self):
        self._usuario('chk_so_conciliacao', 'view_extratoimportado')
        resposta = self._get()
        itens = [e['item'] for e in resposta.context['checklist']]
        self.assertEqual(itens, ['Extrato bancário conciliado'])

    def test_acesso_total_mostra_relatorio_contabil_e_fechamento(self):
        self._usuario(
            'chk_total', 'view_receitaaluguel', 'view_despesa', 'view_documento',
            'view_documentoobrigatorio', 'view_contrato', 'view_extratoimportado', 'view_fechamentomensal',
        )
        resposta = self._get()
        itens = [e['item'] for e in resposta.context['checklist']]
        self.assertIn('Relatório contábil exportado', itens)
        self.assertIn('Fechamento registrado e enviado', itens)

    def test_post_marcar_enviado_continua_protegido_por_change_fechamentomensal(self):
        self._usuario('chk_post_sem_change', 'view_fechamentomensal')
        response = self.client.post(
            reverse('checklist_mensal'),
            {'mes': self.mes, 'ano': self.ano, 'action': 'marcar_enviado'},
        )
        self.assertEqual(response.status_code, 403)


# ─── Testes de baixa de aluguéis filtrada por imóvel ─────────────────────────

class BaixaReceitasFiltroImovelTest(TestCase):
    """Testa filtro por imóvel na tela de baixa de aluguéis."""

    def setUp(self):
        self.client = Client()
        user = com_leitura(User.objects.create_user('filtro_user', password='pass'))
        _dar_permissoes(user, 'add_receitaaluguel', 'change_receitaaluguel')
        self.client.login(username='filtro_user', password='pass')
        locatario = Pessoa.objects.create(nome='Locatário Filtro', tipo='locatario')
        self.imovel1 = Imovel.objects.create(nome='Apto 1', endereco='Rua A', cidade='SP', estado='SP')
        self.imovel2 = Imovel.objects.create(nome='Casa 2', endereco='Rua B', cidade='SP', estado='SP')
        hoje = date.today()
        self.mes, self.ano = hoje.month, hoje.year
        ultimo_dia = monthrange(self.ano, self.mes)[1]
        c1 = _criar_contrato(self.imovel1, locatario, date(self.ano, self.mes, 1), date(self.ano, self.mes, ultimo_dia))
        c2 = _criar_contrato(self.imovel2, locatario, date(self.ano, self.mes, 1), date(self.ano, self.mes, ultimo_dia))
        self.r1 = ReceitaAluguel.objects.create(
            contrato=c1, imovel=self.imovel1,
            competencia_mes=self.mes, competencia_ano=self.ano,
            data_vencimento=date(self.ano, self.mes, 10),
            valor_previsto=Decimal('1000.00'), status='previsto',
        )
        self.r2 = ReceitaAluguel.objects.create(
            contrato=c2, imovel=self.imovel2,
            competencia_mes=self.mes, competencia_ano=self.ano,
            data_vencimento=date(self.ano, self.mes, 10),
            valor_previsto=Decimal('2000.00'), status='previsto',
        )

    def test_get_sem_filtro_retorna_todas_receitas(self):
        """Sem filtro de imóvel, retorna receitas de todos os imóveis."""
        response = self.client.get(reverse('baixa_receitas_mes'), {'mes': self.mes, 'ano': self.ano})
        self.assertEqual(response.status_code, 200)
        self.assertIn(self.r1, response.context['receitas'])
        self.assertIn(self.r2, response.context['receitas'])

    def test_get_com_filtro_retorna_apenas_imovel_selecionado(self):
        """Com filtro de imóvel, retorna apenas as receitas do imóvel."""
        response = self.client.get(
            reverse('baixa_receitas_mes'),
            {'mes': self.mes, 'ano': self.ano, 'imovel': self.imovel1.pk},
        )
        self.assertEqual(response.status_code, 200)
        receitas = list(response.context['receitas'])
        self.assertIn(self.r1, receitas)
        self.assertNotIn(self.r2, receitas)

    def test_post_marcar_recebida_preserva_filtro_imovel(self):
        """POST marcar_recebida com imovel redireciona preservando imovel no URL."""
        response = self.client.post(
            reverse('baixa_receitas_mes'),
            {
                'mes': self.mes, 'ano': self.ano,
                'imovel': self.imovel1.pk,
                'receita_id': self.r1.pk,
                'action': 'marcar_recebida',
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertIn(f'imovel={self.imovel1.pk}', response['Location'])

    def test_post_editar_preserva_filtro_imovel(self):
        """POST editar com imovel redireciona preservando imovel no URL."""
        response = self.client.post(
            reverse('baixa_receitas_mes'),
            {
                'mes': self.mes, 'ano': self.ano,
                'imovel': self.imovel1.pk,
                'receita_id': self.r1.pk,
                'action': 'editar',
                'status': 'parcial',
                'valor_recebido': '800.00',
                'data_recebimento': date.today().isoformat(),
                'observacoes': '',
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertIn(f'imovel={self.imovel1.pk}', response['Location'])

    def test_contexto_contem_imoveis_e_imovel_atual(self):
        """GET com filtro popula imovel_atual e imoveis no contexto."""
        response = self.client.get(
            reverse('baixa_receitas_mes'),
            {'mes': self.mes, 'ano': self.ano, 'imovel': self.imovel1.pk},
        )
        self.assertEqual(str(self.imovel1.pk), response.context['imovel_atual'])
        self.assertIn(self.imovel1, response.context['imoveis'])


# ─── Geração de receitas: prazo indeterminado e encerramento real ────────────

class GerarReceitasPrazoIndeterminadoTest(TestCase):
    """Testa gerar_receitas_mes/gerar_receitas_para_contrato com prazo indeterminado."""

    def setUp(self):
        self.imovel, self.locatario = _criar_base()

    def test_geracao_continua_apos_data_fim_original(self):
        """Contrato por prazo indeterminado gera receita em mês posterior à data_fim original."""
        contrato = _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(2024, 1, 1), data_fim=date(2024, 3, 31),
        )
        contrato.prazo_indeterminado = True
        contrato.save()

        # Gera explicitamente para um mês além da data_fim original (junho/2024)
        criadas, existiam = gerar_receitas_mes(6, 2024)

        self.assertEqual(criadas, 1)
        self.assertTrue(
            ReceitaAluguel.objects.filter(contrato=contrato, competencia_mes=6, competencia_ano=2024).exists()
        )

    def test_data_encerramento_real_limita_geracao(self):
        """Contrato com data_encerramento_real não gera receita após o encerramento."""
        contrato = _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(2024, 1, 1), data_fim=date(2024, 12, 31),
        )
        contrato.prazo_indeterminado = True
        contrato.data_encerramento_real = date(2024, 3, 31)
        contrato.save()

        criadas, existiam = gerar_receitas_mes(6, 2024)

        self.assertEqual(criadas, 0)
        self.assertFalse(
            ReceitaAluguel.objects.filter(contrato=contrato, competencia_mes=6, competencia_ano=2024).exists()
        )

    def test_contrato_encerrado_nao_gera_receita(self):
        """Contrato com status diferente de ativo nunca gera receita, mesmo com prazo indeterminado."""
        contrato = _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(2024, 1, 1), data_fim=date(2024, 3, 31),
            status='encerrado',
        )
        contrato.prazo_indeterminado = True
        contrato.save()

        criadas, existiam = gerar_receitas_mes(6, 2024)
        self.assertEqual(criadas, 0)
        self.assertEqual(existiam, 0)


# ─── Geração de receitas: encargos e itens ────────────────────────────────────

class GerarReceitasComEncargosTest(TestCase):
    """Testa a geração de ReceitaAluguelItem a partir de EncargoContrato."""

    def setUp(self):
        self.imovel, self.locatario = _criar_base()
        self.contrato = _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(2024, 1, 1), data_fim=date(2024, 12, 31),
            valor_aluguel=Decimal('2000.00'),
        )

    def test_contrato_com_iptu_mas_sem_aluguel_gera_aluguel_e_iptu(self):
        """
        Contrato com apenas um encargo de IPTU cadastrado (sem aluguel) deve,
        ao gerar a receita, ter o encargo de aluguel garantido automaticamente
        — a receita final soma aluguel + IPTU, nunca apenas o IPTU.
        """
        EncargoContrato.objects.create(contrato=self.contrato, tipo='iptu', valor=Decimal('150.00'))

        gerar_receitas_para_contrato(self.contrato, data_inicio=date(2024, 3, 1), data_fim=date(2024, 3, 31))

        receita = ReceitaAluguel.objects.get(contrato=self.contrato, competencia_mes=3, competencia_ano=2024)
        tipos = set(receita.itens.values_list('tipo', flat=True))
        self.assertEqual(tipos, {'aluguel', 'iptu'})
        self.assertEqual(receita.valor_previsto, Decimal('2150.00'))

    def test_receita_soma_aluguel_iptu_taxa_manutencao(self):
        EncargoContrato.objects.create(contrato=self.contrato, tipo='aluguel', valor=Decimal('2000.00'))
        EncargoContrato.objects.create(contrato=self.contrato, tipo='iptu', valor=Decimal('150.00'))
        EncargoContrato.objects.create(contrato=self.contrato, tipo='taxa_manutencao', valor=Decimal('300.00'))

        gerar_receitas_para_contrato(self.contrato, data_inicio=date(2024, 3, 1), data_fim=date(2024, 3, 31))

        receita = ReceitaAluguel.objects.get(contrato=self.contrato, competencia_mes=3, competencia_ano=2024)
        self.assertEqual(receita.valor_previsto, Decimal('2450.00'))
        self.assertEqual(receita.itens.count(), 3)
        soma_itens = sum(i.valor for i in receita.itens.all())
        self.assertEqual(soma_itens, receita.valor_previsto)

    def test_encargo_com_inicio_futuro_nao_cobrado_antes(self):
        EncargoContrato.objects.create(contrato=self.contrato, tipo='aluguel', valor=Decimal('2000.00'))
        EncargoContrato.objects.create(
            contrato=self.contrato, tipo='seguro', valor=Decimal('100.00'),
            data_inicio_cobranca=date(2024, 6, 1),
        )

        gerar_receitas_para_contrato(self.contrato, data_inicio=date(2024, 3, 1), data_fim=date(2024, 3, 31))
        gerar_receitas_para_contrato(self.contrato, data_inicio=date(2024, 6, 1), data_fim=date(2024, 6, 30))

        receita_marco = ReceitaAluguel.objects.get(contrato=self.contrato, competencia_mes=3, competencia_ano=2024)
        receita_junho = ReceitaAluguel.objects.get(contrato=self.contrato, competencia_mes=6, competencia_ano=2024)
        self.assertEqual(receita_marco.valor_previsto, Decimal('2000.00'))
        self.assertEqual(receita_junho.valor_previsto, Decimal('2100.00'))

    def test_encargo_encerrado_nao_cobrado_depois(self):
        EncargoContrato.objects.create(contrato=self.contrato, tipo='aluguel', valor=Decimal('2000.00'))
        EncargoContrato.objects.create(
            contrato=self.contrato, tipo='seguro', valor=Decimal('100.00'),
            data_fim_cobranca=date(2024, 3, 31),
        )

        gerar_receitas_para_contrato(self.contrato, data_inicio=date(2024, 3, 1), data_fim=date(2024, 3, 31))
        gerar_receitas_para_contrato(self.contrato, data_inicio=date(2024, 4, 1), data_fim=date(2024, 4, 30))

        receita_marco = ReceitaAluguel.objects.get(contrato=self.contrato, competencia_mes=3, competencia_ano=2024)
        receita_abril = ReceitaAluguel.objects.get(contrato=self.contrato, competencia_mes=4, competencia_ano=2024)
        self.assertEqual(receita_marco.valor_previsto, Decimal('2100.00'))
        self.assertEqual(receita_abril.valor_previsto, Decimal('2000.00'))

    def test_fallback_valor_aluguel_sem_encargos(self):
        """
        Contrato sem nenhum EncargoContrato continua usando valor_aluguel — a
        geração garante automaticamente um encargo de aluguel (item 1 desta
        fase), então a receita passa a ter exatamente 1 item (aluguel).
        """
        gerar_receitas_para_contrato(self.contrato, data_inicio=date(2024, 3, 1), data_fim=date(2024, 3, 31))

        receita = ReceitaAluguel.objects.get(contrato=self.contrato, competencia_mes=3, competencia_ano=2024)
        self.assertEqual(receita.valor_previsto, self.contrato.valor_aluguel)
        self.assertEqual(receita.itens.count(), 1)
        self.assertEqual(receita.itens.first().tipo, 'aluguel')

    def test_receitas_ja_existentes_nao_sao_alteradas(self):
        """Gerar novamente não sobrescreve receita/itens já existentes."""
        EncargoContrato.objects.create(contrato=self.contrato, tipo='aluguel', valor=Decimal('2000.00'))
        gerar_receitas_para_contrato(self.contrato, data_inicio=date(2024, 3, 1), data_fim=date(2024, 3, 31))

        receita = ReceitaAluguel.objects.get(contrato=self.contrato, competencia_mes=3, competencia_ano=2024)
        receita.valor_recebido = Decimal('2000.00')
        receita.status = 'recebido'
        receita.save()

        # Adiciona um novo encargo depois — não deve afetar a receita já criada
        EncargoContrato.objects.create(contrato=self.contrato, tipo='iptu', valor=Decimal('150.00'))
        criadas, existiam = gerar_receitas_para_contrato(
            self.contrato, data_inicio=date(2024, 3, 1), data_fim=date(2024, 3, 31)
        )

        self.assertEqual(criadas, 0)
        self.assertEqual(existiam, 1)
        receita.refresh_from_db()
        self.assertEqual(receita.valor_previsto, Decimal('2000.00'))
        self.assertEqual(receita.status, 'recebido')
        # O novo encargo IPTU (criado depois) não deve retroagir sobre a receita já existente
        self.assertEqual(receita.itens.count(), 1)


# ─── Taxa de administração da imobiliária (despesa automática) ────────────────

class TaxaAdministracaoDespesaTest(TestCase):
    def setUp(self):
        self.imovel, self.locatario = _criar_base()

    def test_contrato_com_taxa_gera_despesa_automatica(self):
        contrato = _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(2024, 1, 1), data_fim=date(2024, 12, 31),
            valor_aluguel=Decimal('2000.00'),
        )
        contrato.comissao_imobiliaria_percentual = Decimal('10.00')
        contrato.save()

        gerar_receitas_para_contrato(contrato, data_inicio=date(2024, 3, 1), data_fim=date(2024, 3, 31))

        despesa = Despesa.objects.get(contrato=contrato, origem_automatica=True, categoria='comissao_imobiliaria')
        self.assertEqual(despesa.valor, Decimal('200.00'))

    def test_contrato_sem_taxa_nao_gera_despesa(self):
        contrato = _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(2024, 1, 1), data_fim=date(2024, 12, 31),
        )
        gerar_receitas_para_contrato(contrato, data_inicio=date(2024, 3, 1), data_fim=date(2024, 3, 31))
        self.assertFalse(Despesa.objects.filter(contrato=contrato, origem_automatica=True).exists())

    def test_gerar_duas_vezes_nao_duplica_despesa(self):
        contrato = _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(2024, 1, 1), data_fim=date(2024, 12, 31),
        )
        contrato.comissao_imobiliaria_percentual = Decimal('8.00')
        contrato.save()

        gerar_receitas_para_contrato(contrato, data_inicio=date(2024, 3, 1), data_fim=date(2024, 3, 31))
        gerar_receitas_para_contrato(contrato, data_inicio=date(2024, 3, 1), data_fim=date(2024, 3, 31))

        self.assertEqual(
            Despesa.objects.filter(contrato=contrato, origem_automatica=True).count(), 1
        )

    def test_despesa_automatica_tem_fornecedor_igual_a_imobiliaria_do_contrato(self):
        imobiliaria = Pessoa.objects.create(nome='Imob Padrão', tipo='imobiliaria')
        contrato = _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(2024, 1, 1), data_fim=date(2024, 12, 31),
            valor_aluguel=Decimal('2000.00'), imobiliaria=imobiliaria,
        )
        contrato.comissao_imobiliaria_percentual = Decimal('10.00')
        contrato.save()

        gerar_receitas_para_contrato(contrato, data_inicio=date(2024, 3, 1), data_fim=date(2024, 3, 31))

        despesa = Despesa.objects.get(contrato=contrato, origem_automatica=True, categoria='comissao_imobiliaria')
        self.assertEqual(despesa.fornecedor_id, imobiliaria.pk)

    def test_despesa_automatica_sem_imobiliaria_no_contrato_fica_sem_fornecedor(self):
        contrato = _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(2024, 1, 1), data_fim=date(2024, 12, 31),
            valor_aluguel=Decimal('2000.00'),
        )
        contrato.comissao_imobiliaria_percentual = Decimal('10.00')
        contrato.save()

        gerar_receitas_para_contrato(contrato, data_inicio=date(2024, 3, 1), data_fim=date(2024, 3, 31))

        despesa = Despesa.objects.get(contrato=contrato, origem_automatica=True, categoria='comissao_imobiliaria')
        self.assertIsNone(despesa.fornecedor_id)

    def test_taxa_calculada_sobre_encargo_aluguel_nao_sobre_total(self):
        contrato = _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(2024, 1, 1), data_fim=date(2024, 12, 31),
            valor_aluguel=Decimal('2000.00'),
        )
        contrato.comissao_imobiliaria_percentual = Decimal('10.00')
        contrato.save()
        EncargoContrato.objects.create(contrato=contrato, tipo='aluguel', valor=Decimal('2000.00'))
        EncargoContrato.objects.create(contrato=contrato, tipo='iptu', valor=Decimal('500.00'))

        gerar_receitas_para_contrato(contrato, data_inicio=date(2024, 3, 1), data_fim=date(2024, 3, 31))

        despesa = Despesa.objects.get(contrato=contrato, origem_automatica=True)
        # 10% sobre 2000 (aluguel) — não sobre 2500 (aluguel + IPTU)
        self.assertEqual(despesa.valor, Decimal('200.00'))


# ─── Multa e juros por atraso ──────────────────────────────────────────────────

class CalcularMultaJurosTest(TestCase):
    def setUp(self):
        self.imovel, self.locatario = _criar_base()
        self.contrato = _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(2020, 1, 1), data_fim=date(2030, 12, 31),
        )
        self.contrato.multa_atraso_percentual = Decimal('2.00')
        self.contrato.juros_mora_percentual_mes = Decimal('1.00')
        self.contrato.dias_carencia_multa = 3
        self.contrato.save()

    def _receita(self, vencimento):
        return ReceitaAluguel(
            contrato=self.contrato, imovel=self.imovel,
            competencia_mes=1, competencia_ano=2024,
            data_vencimento=vencimento, valor_previsto=Decimal('1000.00'),
            status='previsto',
        )

    def test_sem_atraso_multa_juros_zero(self):
        receita = self._receita(date.today() + timedelta(days=1))
        resultado = receita.calcular_multa_juros(date.today())
        self.assertEqual(resultado['multa'], Decimal('0.00'))
        self.assertEqual(resultado['juros'], Decimal('0.00'))

    def test_atraso_alem_da_carencia_gera_multa_e_juros(self):
        vencimento = date(2024, 1, 1)
        data_recebimento = date(2024, 1, 15)  # 14 dias de atraso, carência de 3
        receita = self._receita(vencimento)
        resultado = receita.calcular_multa_juros(data_recebimento)

        self.assertEqual(resultado['dias_atraso'], 14)
        self.assertEqual(resultado['multa'], Decimal('20.00'))  # 2% de 1000
        # juros: 1000 * 1% / 30 * 14 = 4.67
        self.assertEqual(resultado['juros'], Decimal('4.67'))

    def test_atraso_dentro_da_carencia_nao_gera_multa(self):
        vencimento = date(2024, 1, 1)
        data_recebimento = date(2024, 1, 3)  # 2 dias de atraso, carência de 3
        receita = self._receita(vencimento)
        resultado = receita.calcular_multa_juros(data_recebimento)
        self.assertEqual(resultado['multa'], Decimal('0.00'))
        self.assertEqual(resultado['juros'], Decimal('0.00'))

    def test_calculo_nao_altera_a_receita(self):
        """calcular_multa_juros apenas retorna a sugestão — nunca salva no banco."""
        receita = ReceitaAluguel.objects.create(
            contrato=self.contrato, imovel=self.imovel,
            competencia_mes=1, competencia_ano=2024,
            data_vencimento=date(2024, 1, 1), valor_previsto=Decimal('1000.00'),
            multa=Decimal('5.00'), juros=Decimal('3.00'), status='previsto',
        )
        receita.calcular_multa_juros(date(2024, 1, 20))
        receita.refresh_from_db()
        self.assertEqual(receita.multa, Decimal('5.00'))
        self.assertEqual(receita.juros, Decimal('3.00'))


# ─── Checklist: etapa de reajustes pendentes ──────────────────────────────────

class ChecklistReajustesPendentesTest(TestCase):
    def setUp(self):
        self.client = Client()
        com_leitura(User.objects.create_user('chk_reaj_user', password='pass'))
        self.client.login(username='chk_reaj_user', password='pass')
        self.imovel, self.locatario = _criar_base()

    def test_etapa_reajustes_nao_ok_quando_pendente(self):
        contrato = _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(2020, 1, 1), data_fim=date(2030, 12, 31),
        )
        contrato.data_proximo_reajuste = date.today() - timedelta(days=1)
        contrato.save()

        response = self.client.get(reverse('checklist_mensal'))
        checklist = response.context['checklist']
        etapa = next(e for e in checklist if e['item'] == 'Reajustes em dia')
        self.assertFalse(etapa['ok'])

    def test_etapa_reajustes_ok_quando_sem_pendencia(self):
        response = self.client.get(reverse('checklist_mensal'))
        checklist = response.context['checklist']
        etapa = next(e for e in checklist if e['item'] == 'Reajustes em dia')
        self.assertTrue(etapa['ok'])


# ─── Teste de migrations pendentes ────────────────────────────────────────────

class MigracoesPendentesTest(TestCase):
    """Verifica que não há migrations pendentes após todas as alterações."""

    def test_nao_ha_migrations_pendentes(self):
        from io import StringIO
        from django.core.management import call_command
        try:
            call_command('makemigrations', '--check', verbosity=0)
        except SystemExit as e:
            self.fail('Há migrations pendentes não criadas.')


# ─── Parâmetros inválidos na URL não podem derrubar as views ──────────────────

class ParametrosInvalidosViewTest(TestCase):
    def setUp(self):
        self.client = Client()
        com_leitura(User.objects.create_user('param_user', password='pass'))
        self.client.login(username='param_user', password='pass')

    def test_receitas_list_mes_nao_numerico_usa_padrao(self):
        response = self.client.get(reverse('receitas_list'), {'mes': 'abc', 'ano': ''})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['mes_atual'], timezone.localdate().month)

    def test_receitas_list_mes_fora_da_faixa_usa_padrao(self):
        response = self.client.get(reverse('receitas_list'), {'mes': '13', 'ano': '2024'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['mes_atual'], timezone.localdate().month)

    def test_receitas_list_imovel_nao_numerico_ignorado(self):
        response = self.client.get(reverse('receitas_list'), {'imovel': 'abc'})
        self.assertEqual(response.status_code, 200)

    def test_despesas_list_parametros_invalidos_nao_quebram(self):
        response = self.client.get(reverse('despesas_list'), {'mes': 'x', 'ano': 'y', 'imovel': 'z'})
        self.assertEqual(response.status_code, 200)

    def test_export_receitas_parametros_invalidos_nao_quebram(self):
        response = self.client.get(reverse('export_receitas', args=['csv']), {'mes': 'abc', 'imovel': 'abc'})
        self.assertEqual(response.status_code, 200)

    def test_baixa_receitas_parametros_invalidos_nao_quebram(self):
        response = self.client.get(reverse('baixa_receitas_mes'), {'mes': '99', 'ano': 'abc'})
        self.assertEqual(response.status_code, 200)


# ─── Validação da edição na baixa de receitas ─────────────────────────────────

class BaixaReceitasValidacaoTest(TestCase):
    def setUp(self):
        self.client = Client()
        self.user = com_leitura(User.objects.create_user('valida_user', password='pass'))
        _dar_permissoes(self.user, 'add_receitaaluguel', 'change_receitaaluguel')
        self.client.login(username='valida_user', password='pass')
        self.imovel, self.locatario = _criar_base()
        hoje = date.today()
        self.mes, self.ano = hoje.month, hoje.year
        ultimo_dia = monthrange(self.ano, self.mes)[1]
        self.contrato = _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(self.ano, self.mes, 1),
            data_fim=date(self.ano, self.mes, ultimo_dia),
        )
        self.receita = ReceitaAluguel.objects.create(
            contrato=self.contrato, imovel=self.imovel,
            competencia_mes=self.mes, competencia_ano=self.ano,
            data_vencimento=date(self.ano, self.mes, 10),
            valor_previsto=Decimal('2000.00'), status='previsto',
        )

    def _post_recebimento(self, **campos):
        dados = {
            'mes': self.mes, 'ano': self.ano,
            'receita_id': self.receita.pk, 'action': 'registrar_recebimento',
        }
        dados.update(campos)
        return self.client.post(reverse('baixa_receitas_mes'), dados)

    def _post_encargos(self, **campos):
        dados = {
            'mes': self.mes, 'ano': self.ano,
            'receita_id': self.receita.pk, 'action': 'editar_encargos',
        }
        dados.update(campos)
        return self.client.post(reverse('baixa_receitas_mes'), dados)

    def test_valor_invalido_nao_altera_receita(self):
        response = self._post_recebimento(valor='abc', data_recebimento=date.today().isoformat())
        self.assertEqual(response.status_code, 302)
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.status, 'previsto')
        self.assertIsNone(self.receita.valor_recebido)
        self.assertEqual(self.receita.recebimentos.count(), 0)

    def test_data_invalida_nao_altera_receita(self):
        response = self._post_recebimento(valor='2000.00', data_recebimento='31/31/2024')
        self.assertEqual(response.status_code, 302)
        self.receita.refresh_from_db()
        self.assertIsNone(self.receita.data_recebimento)
        self.assertEqual(self.receita.recebimentos.count(), 0)

    def test_valor_zero_rejeitado(self):
        self._post_recebimento(valor='0', data_recebimento=date.today().isoformat())
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.recebimentos.count(), 0)

    def test_valor_acima_do_saldo_rejeitado(self):
        self._post_recebimento(valor='2000.01', data_recebimento=date.today().isoformat())
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.recebimentos.count(), 0)
        self.assertIsNone(self.receita.valor_recebido)

    def test_edicao_de_encargos_atualiza_e_reconsolida(self):
        # vencimento futuro para o status reconsolidado ser determinístico
        self.receita.data_vencimento = date.today() + timedelta(days=5)
        self.receita.save(update_fields=['data_vencimento'])
        self._post_encargos(multa='10.00', juros='5.25', desconto='2.00', observacoes='ajuste')
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.multa, Decimal('10.00'))
        self.assertEqual(self.receita.juros, Decimal('5.25'))
        self.assertEqual(self.receita.desconto, Decimal('2.00'))
        self.assertEqual(self.receita.observacoes, 'ajuste')
        # consolidado permanece derivado dos recebimentos (nenhum registrado)
        self.assertIsNone(self.receita.valor_recebido)
        self.assertEqual(self.receita.status, 'previsto')

    def test_desconto_maior_que_devido_rejeitado(self):
        self._post_encargos(desconto='2500.00')
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.desconto, Decimal('0.00'))

    def test_multa_apos_pagamento_parcial_recalcula_status(self):
        # paga integralmente (2000) → recebido
        self._post_recebimento(valor='2000.00', data_recebimento=date.today().isoformat())
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.status, 'recebido')
        # multa de 100 reabre saldo de 100 → volta a parcial
        self._post_encargos(multa='100.00')
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.status, 'parcial')
        self.assertEqual(self.receita.saldo_em_aberto, Decimal('100.00'))

    def test_cancelar_e_reabrir_receita(self):
        self._post_recebimento(valor='500.00', data_recebimento=date.today().isoformat())
        self.client.post(reverse('baixa_receitas_mes'), {
            'mes': self.mes, 'ano': self.ano,
            'receita_id': self.receita.pk, 'action': 'cancelar',
        })
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.status, 'cancelado')
        # cancelamento encerra a cobrança sem apagar recebimentos
        self.assertEqual(self.receita.saldo_em_aberto, Decimal('0.00'))
        self.assertTrue(self.receita.esta_quitada)
        self.assertEqual(self.receita.recebimentos.count(), 1)
        # reabrir recalcula pelo saldo/recebimentos → parcial
        self.client.post(reverse('baixa_receitas_mes'), {
            'mes': self.mes, 'ano': self.ano,
            'receita_id': self.receita.pk, 'action': 'reabrir',
        })
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.status, 'parcial')
        self.assertEqual(self.receita.valor_recebido, Decimal('500.00'))

    def test_marcar_recebida_quita_o_saldo_integral(self):
        """marcar_recebida registra o recebimento do saldo em aberto integral."""
        self.receita.valor_recebido = Decimal('0.00')
        self.receita.save()
        self.client.post(reverse('baixa_receitas_mes'), {
            'mes': self.mes, 'ano': self.ano,
            'receita_id': self.receita.pk, 'action': 'marcar_recebida',
        })
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.valor_recebido, Decimal('2000.00'))
        self.assertEqual(self.receita.status, 'recebido')
        self.assertEqual(self.receita.saldo_em_aberto, Decimal('0.00'))

    def test_post_sem_permissao_retorna_403(self):
        com_leitura(User.objects.create_user('sem_perm', password='pass'))
        self.client.login(username='sem_perm', password='pass')
        response = self.client.post(reverse('baixa_receitas_mes'), {
            'mes': self.mes, 'ano': self.ano,
            'receita_id': self.receita.pk, 'action': 'marcar_recebida',
        })
        self.assertEqual(response.status_code, 403)
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.status, 'previsto')

    def test_gerar_receitas_sem_permissao_retorna_403(self):
        com_leitura(User.objects.create_user('sem_perm2', password='pass'))
        self.client.login(username='sem_perm2', password='pass')
        response = self.client.post(reverse('gerar_receitas_mes'), {'mes': self.mes, 'ano': self.ano})
        self.assertEqual(response.status_code, 403)


# ─── Comissão de imobiliária respeita a carência do encargo de aluguel ────────

class ComissaoCarenciaTest(TestCase):
    def setUp(self):
        self.imovel, self.locatario = _criar_base()
        self.contrato = _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(2024, 1, 1), data_fim=date(2024, 12, 31),
            comissao_imobiliaria_percentual=Decimal('10.00'),
        )
        EncargoContrato.objects.create(
            contrato=self.contrato, tipo='aluguel', valor=Decimal('2000.00'),
            data_inicio_cobranca=date(2024, 6, 1),
        )

    def test_mes_de_carencia_nao_gera_comissao(self):
        gerar_receitas_para_contrato(self.contrato, data_inicio=date(2024, 3, 1), data_fim=date(2024, 3, 31))
        self.assertFalse(
            Despesa.objects.filter(contrato=self.contrato, categoria='comissao_imobiliaria').exists()
        )

    def test_mes_com_aluguel_gera_comissao(self):
        gerar_receitas_para_contrato(self.contrato, data_inicio=date(2024, 6, 1), data_fim=date(2024, 6, 30))
        despesa = Despesa.objects.get(contrato=self.contrato, categoria='comissao_imobiliaria')
        self.assertEqual(despesa.valor, Decimal('200.00'))


# ─── Comando marcar_atrasados ─────────────────────────────────────────────────

class MarcarAtrasadosCommandTest(TestCase):
    def setUp(self):
        self.imovel, self.locatario = _criar_base()
        self.contrato = _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(2020, 1, 1), data_fim=date(2030, 12, 31),
        )
        self.ontem = timezone.localdate() - timedelta(days=1)
        self.amanha = timezone.localdate() + timedelta(days=1)

    def test_comando_marca_receitas_e_despesas_vencidas(self):
        from django.core.management import call_command

        vencida = ReceitaAluguel.objects.create(
            contrato=self.contrato, imovel=self.imovel,
            competencia_mes=1, competencia_ano=2024,
            data_vencimento=self.ontem, valor_previsto=Decimal('1000.00'), status='previsto',
        )
        futura = ReceitaAluguel.objects.create(
            contrato=self.contrato, imovel=self.imovel,
            competencia_mes=2, competencia_ano=2024,
            data_vencimento=self.amanha, valor_previsto=Decimal('1000.00'), status='previsto',
        )
        recebida = ReceitaAluguel.objects.create(
            contrato=self.contrato, imovel=self.imovel,
            competencia_mes=3, competencia_ano=2024,
            data_vencimento=self.ontem, valor_previsto=Decimal('1000.00'), status='recebido',
        )
        despesa_vencida = Despesa.objects.create(
            imovel=self.imovel, descricao='IPTU', data_vencimento=self.ontem,
            valor=Decimal('500.00'), status='prevista',
        )

        call_command('marcar_atrasados')

        vencida.refresh_from_db()
        futura.refresh_from_db()
        recebida.refresh_from_db()
        despesa_vencida.refresh_from_db()
        self.assertEqual(vencida.status, 'atrasado')
        self.assertEqual(futura.status, 'previsto')
        self.assertEqual(recebida.status, 'recebido')
        self.assertEqual(despesa_vencida.status, 'atrasada')


# ─── Série de fluxo de caixa de 12 meses ──────────────────────────────────────

class FluxoCaixa12mTest(TestCase):
    def test_serie_cobre_12_meses_e_soma_por_data_de_pagamento(self):
        from financeiro.services import serie_fluxo_caixa_12m, registrar_recebimento

        imovel, locatario = _criar_base()
        contrato = _criar_contrato(
            imovel, locatario, data_inicio=date(2020, 1, 1), data_fim=date(2030, 12, 31)
        )
        hoje = timezone.localdate()
        receita = ReceitaAluguel.objects.create(
            contrato=contrato, imovel=imovel,
            competencia_mes=hoje.month, competencia_ano=hoje.year,
            data_vencimento=hoje, valor_previsto=Decimal('2000.00'), status='previsto',
        )
        registrar_recebimento(receita.pk, Decimal('2000.00'), hoje)
        Despesa.objects.create(
            imovel=imovel, descricao='Condomínio',
            competencia_mes=hoje.month, competencia_ano=hoje.year,
            data_vencimento=hoje, data_pagamento=hoje, valor=Decimal('300.00'), status='paga',
        )

        serie = serie_fluxo_caixa_12m(hoje)
        self.assertEqual(len(serie), 12)
        atual = serie[-1]
        self.assertEqual(atual['label'], f'{hoje.month:02d}/{hoje.year}')
        self.assertEqual(atual['recebido'], 2000.0)
        self.assertEqual(atual['pago'], 300.0)
        self.assertEqual(atual['saldo'], 1700.0)
        # Receita prevista (não recebida) não entra na série
        self.assertEqual(sum(p['recebido'] for p in serie), 2000.0)

    def test_aluguel_de_janeiro_pago_em_marco_aparece_em_marco(self):
        """Item 5: fluxo de caixa usa a data real do pagamento, não a competência."""
        from financeiro.services import serie_fluxo_caixa_12m, registrar_recebimento

        imovel, locatario = _criar_base()
        contrato = _criar_contrato(
            imovel, locatario, data_inicio=date(2020, 1, 1), data_fim=date(2030, 12, 31)
        )
        receita = ReceitaAluguel.objects.create(
            contrato=contrato, imovel=imovel,
            competencia_mes=1, competencia_ano=2030,
            data_vencimento=date(2030, 1, 10), valor_previsto=Decimal('1500.00'), status='previsto',
        )
        # Pago em março, não em janeiro
        registrar_recebimento(receita.pk, Decimal('1500.00'), date(2030, 3, 5))

        serie = serie_fluxo_caixa_12m(date(2030, 3, 1))
        por_label = {p['label']: p for p in serie}
        self.assertEqual(por_label['03/2030']['recebido'], 1500.0)
        self.assertEqual(por_label['01/2030']['recebido'], 0.0)

    def test_receita_cancelada_com_pagamento_continua_no_fluxo_de_caixa(self):
        """Item 4: dinheiro recebido continua no caixa mesmo após cancelamento."""
        from financeiro.services import serie_fluxo_caixa_12m, registrar_recebimento

        imovel, locatario = _criar_base()
        contrato = _criar_contrato(
            imovel, locatario, data_inicio=date(2020, 1, 1), data_fim=date(2030, 12, 31)
        )
        hoje = timezone.localdate()
        receita = ReceitaAluguel.objects.create(
            contrato=contrato, imovel=imovel,
            competencia_mes=hoje.month, competencia_ano=hoje.year,
            data_vencimento=hoje, valor_previsto=Decimal('2000.00'), status='previsto',
        )
        registrar_recebimento(receita.pk, Decimal('500.00'), hoje)
        receita.cancelar()

        serie = serie_fluxo_caixa_12m(hoje)
        self.assertEqual(sum(p['recebido'] for p in serie), 500.0)

    def test_incluir_despesas_false_omite_despesas_pagas(self):
        from financeiro.services import serie_fluxo_caixa_12m

        imovel, locatario = _criar_base()
        contrato = _criar_contrato(
            imovel, locatario, data_inicio=date(2020, 1, 1), data_fim=date(2030, 12, 31)
        )
        hoje = timezone.localdate()
        Despesa.objects.create(
            imovel=imovel, descricao='Condomínio',
            competencia_mes=hoje.month, competencia_ano=hoje.year,
            data_vencimento=hoje, data_pagamento=hoje, valor=Decimal('300.00'), status='paga',
        )
        serie = serie_fluxo_caixa_12m(hoje, incluir_despesas=False)
        # Item 9: sem despesas incluídas, 'pago'/'saldo' são None (não 0.0) —
        # None sinaliza "não consultado", nunca "despesa zero".
        self.assertTrue(all(p['pago'] is None for p in serie))
        self.assertTrue(all(p['saldo'] is None for p in serie))


# ─── Painéis gráficos ─────────────────────────────────────────────────────────

class PaineisSeriesTest(TestCase):
    def setUp(self):
        self.imovel, self.locatario = _criar_base()
        self.hoje = timezone.localdate()

    def test_serie_ocupacao_conta_contrato_vigente(self):
        from financeiro.paineis import serie_ocupacao_mensal

        # Imóvel agrupador não entra na base de locáveis
        Imovel.objects.create(
            nome='Prédio Agrupador', endereco='X', cidade='SP', estado='SP', unidade_locavel=False,
        )
        _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=self.hoje - timedelta(days=400),
            data_fim=self.hoje + timedelta(days=365),
        )
        vago = Imovel.objects.create(nome='Imóvel Vago', endereco='Y', cidade='SP', estado='SP')

        serie = serie_ocupacao_mensal(12, self.hoje)
        self.assertEqual(len(serie), 12)
        atual = serie[-1]
        self.assertEqual(atual['total'], 2)          # ocupado + vago (agrupador fora)
        self.assertEqual(atual['ocupados'], 1)
        self.assertEqual(atual['percentual'], 50.0)
        self.assertEqual(atual['vacancia'], 50.0)

    def test_serie_ocupacao_respeita_encerramento_real(self):
        from financeiro.paineis import serie_ocupacao_mensal

        _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=self.hoje - timedelta(days=400),
            data_fim=self.hoje + timedelta(days=365),
            status='encerrado',
            data_encerramento_real=self.hoje - timedelta(days=200),
        )
        serie = serie_ocupacao_mensal(3, self.hoje)
        self.assertEqual(serie[-1]['ocupados'], 0)

    def test_serie_receita_despesa_por_imovel(self):
        from financeiro.paineis import serie_receita_despesa_por_imovel
        from financeiro.services import registrar_recebimento

        contrato = _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(2020, 1, 1), data_fim=date(2030, 12, 31),
        )
        receita = ReceitaAluguel.objects.create(
            contrato=contrato, imovel=self.imovel,
            competencia_mes=self.hoje.month, competencia_ano=self.hoje.year,
            data_vencimento=self.hoje, valor_previsto=Decimal('2000.00'), status='previsto',
        )
        registrar_recebimento(receita.pk, Decimal('2000.00'), self.hoje)
        Despesa.objects.create(
            imovel=self.imovel, descricao='Condomínio',
            competencia_mes=self.hoje.month, competencia_ano=self.hoje.year,
            data_vencimento=self.hoje, data_pagamento=self.hoje, valor=Decimal('300.00'), status='paga',
        )
        serie = serie_receita_despesa_por_imovel(12, self.hoje)
        self.assertEqual(len(serie), 1)
        self.assertEqual(serie[0]['imovel'], self.imovel.nome)
        self.assertEqual(serie[0]['recebido'], 2000.0)
        self.assertEqual(serie[0]['pago'], 300.0)
        self.assertEqual(serie[0]['resultado'], 1700.0)

    def test_serie_inadimplencia_mensal(self):
        from financeiro.paineis import serie_inadimplencia_mensal

        contrato = _criar_contrato(
            self.imovel, self.locatario,
            data_inicio=date(2020, 1, 1), data_fim=date(2030, 12, 31),
        )
        ontem = self.hoje - timedelta(days=1)
        ReceitaAluguel.objects.create(
            contrato=contrato, imovel=self.imovel,
            competencia_mes=ontem.month, competencia_ano=ontem.year,
            data_vencimento=ontem, valor_previsto=Decimal('1500.00'), status='previsto',
        )
        serie = serie_inadimplencia_mensal(12, self.hoje)
        self.assertEqual(len(serie), 12)
        alvo = next(p for p in serie if p['label'] == f'{ontem.month:02d}/{ontem.year}')
        self.assertEqual(alvo['valor'], 1500.0)
        self.assertEqual(alvo['quantidade'], 1)

    def test_serie_despesas_por_categoria_exclui_canceladas(self):
        from financeiro.paineis import serie_despesas_por_categoria

        Despesa.objects.create(
            imovel=self.imovel, descricao='IPTU', categoria='iptu',
            competencia_mes=self.hoje.month, competencia_ano=self.hoje.year,
            data_vencimento=self.hoje, data_pagamento=self.hoje, valor=Decimal('800.00'), status='paga',
        )
        Despesa.objects.create(
            imovel=self.imovel, descricao='Cancelada', categoria='manutencao',
            competencia_mes=self.hoje.month, competencia_ano=self.hoje.year,
            data_vencimento=self.hoje, valor=Decimal('999.00'), status='cancelada',
        )
        serie = serie_despesas_por_categoria(12, self.hoje)
        self.assertEqual(len(serie), 1)
        self.assertEqual(serie[0]['categoria'], 'IPTU')
        self.assertEqual(serie[0]['valor'], 800.0)


class PaineisViewTest(TestCase):
    def setUp(self):
        self.client = Client()
        com_leitura(User.objects.create_user('painel_user', password='pass'))
        self.client.login(username='painel_user', password='pass')

    def test_exige_login(self):
        self.client.logout()
        response = self.client.get(reverse('paineis'))
        self.assertEqual(response.status_code, 302)
        self.assertIn('/login/', response['Location'])

    def test_renderiza_com_series(self):
        response = self.client.get(reverse('paineis'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.context['ocupacao']), 12)
        self.assertEqual(len(response.context['inadimplencia']), 12)
        self.assertContains(response, 'grafico-ocupacao')
        self.assertContains(response, 'grafico-inadimplencia')
        self.assertContains(response, 'grafico-por-imovel')
        self.assertContains(response, 'grafico-categorias')

    def test_janela_invalida_usa_padrao(self):
        response = self.client.get(reverse('paineis'), {'janela': '99', 'mes': 'abc'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['janela_atual'], 12)

    def test_janela_24_meses(self):
        response = self.client.get(reverse('paineis'), {'janela': '24'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.context['ocupacao']), 24)


class PaineisPermissoesTest(TestCase):
    """Item 3 (rodada pós-revisão): painéis escondem despesas/resultado sem permissão."""

    def setUp(self):
        from django.contrib.auth.models import Permission
        self.Permission = Permission
        self.client = Client()
        imovel, locatario = _criar_base()
        contrato = _criar_contrato(imovel, locatario, date(2024, 1, 1), date(2024, 12, 31))
        Despesa.objects.create(
            imovel=imovel, categoria='iptu', descricao='IPTU', data_vencimento=date(2024, 3, 10),
            data_pagamento=date(2024, 3, 10), valor=Decimal('200.00'), status='paga',
        )

    def _usuario(self, nome, *perms):
        user = User.objects.create_user(nome, password='pass')
        if perms:
            user.user_permissions.add(*self.Permission.objects.filter(codename__in=perms))
        self.client.login(username=nome, password='pass')
        return user

    def _get(self):
        return self.client.get(reverse('paineis'))

    def test_sem_nenhuma_permissao_403(self):
        self._usuario('painel_sem_perm')
        self.assertEqual(self._get().status_code, 403)

    def test_apenas_receitas_esconde_despesas_e_resultado(self):
        self._usuario('painel_so_receitas', 'view_receitaaluguel')
        resposta = self._get()
        self.assertEqual(resposta.status_code, 200)
        conteudo = resposta.content.decode()
        self.assertIsNone(resposta.context['por_categoria'])
        self.assertIsNone(resposta.context['ocupacao'])
        self.assertIsNotNone(resposta.context['inadimplencia'])
        self.assertIsNotNone(resposta.context['por_imovel'])
        for linha in resposta.context['por_imovel']:
            self.assertIsNone(linha['pago'])
            self.assertIsNone(linha['resultado'])
        self.assertNotIn('id="grafico-categorias"', conteudo)
        self.assertNotIn('id="grafico-ocupacao"', conteudo)
        self.assertIn('id="grafico-por-imovel"', conteudo)
        self.assertIn('Receitas recebidas por imóvel', conteudo)
        self.assertNotIn('Pago (R$)', conteudo)

    def test_apenas_despesas_esconde_receitas_e_inadimplencia(self):
        self._usuario('painel_so_despesas', 'view_despesa')
        resposta = self._get()
        self.assertEqual(resposta.status_code, 200)
        conteudo = resposta.content.decode()
        self.assertIsNone(resposta.context['inadimplencia'])
        self.assertIsNotNone(resposta.context['por_categoria'])
        self.assertIsNotNone(resposta.context['por_imovel'])
        for linha in resposta.context['por_imovel']:
            self.assertIsNone(linha['recebido'])
            self.assertIsNone(linha['resultado'])
        self.assertNotIn('id="grafico-inadimplencia"', conteudo)
        self.assertIn('id="grafico-categorias"', conteudo)
        self.assertIn('Despesas pagas por imóvel', conteudo)
        self.assertNotIn('Recebido (R$)', conteudo)

    def test_apenas_patrimonio_mostra_so_ocupacao(self):
        self._usuario('painel_so_patrimonio', 'view_imovel')
        resposta = self._get()
        self.assertEqual(resposta.status_code, 200)
        conteudo = resposta.content.decode()
        self.assertIsNotNone(resposta.context['ocupacao'])
        self.assertIsNone(resposta.context['inadimplencia'])
        self.assertIsNone(resposta.context['por_categoria'])
        self.assertIsNone(resposta.context['por_imovel'])
        self.assertIn('id="grafico-ocupacao"', conteudo)
        self.assertNotIn('id="grafico-por-imovel"', conteudo)
        self.assertNotIn('id="grafico-categorias"', conteudo)
        self.assertNotIn('id="grafico-inadimplencia"', conteudo)

    def test_receitas_e_despesas_mostra_tudo_com_resultado(self):
        self._usuario('painel_completo', 'view_receitaaluguel', 'view_despesa')
        resposta = self._get()
        self.assertEqual(resposta.status_code, 200)
        conteudo = resposta.content.decode()
        self.assertIsNotNone(resposta.context['por_categoria'])
        self.assertIsNotNone(resposta.context['inadimplencia'])
        for linha in resposta.context['por_imovel']:
            self.assertIsNotNone(linha['recebido'])
            self.assertIsNotNone(linha['pago'])
            self.assertIsNotNone(linha['resultado'])
        self.assertIn('Receitas × despesas por imóvel', conteudo)
        self.assertIn('Resultado (R$)', conteudo)


# ─── Pagamentos parciais e RecebimentoReceita (rodada de correções) ───────────

class PagamentosParciaisTest(TestCase):
    def setUp(self):
        from financeiro.models import RecebimentoReceita
        self.RecebimentoReceita = RecebimentoReceita
        self.imovel, self.locatario = _criar_base()
        self.contrato = _criar_contrato(
            self.imovel, self.locatario, data_inicio=date(2020, 1, 1), data_fim=date(2030, 12, 31),
        )
        self.ontem = timezone.localdate() - timedelta(days=1)
        self.receita = ReceitaAluguel.objects.create(
            contrato=self.contrato, imovel=self.imovel,
            competencia_mes=self.ontem.month, competencia_ano=self.ontem.year,
            data_vencimento=self.ontem, valor_previsto=Decimal('2000.00'), status='previsto',
        )

    def _receber(self, valor):
        return self.RecebimentoReceita.objects.create(
            receita=self.receita, data_recebimento=timezone.localdate(), valor=Decimal(valor),
        )

    def test_parcial_continua_em_aberto_e_na_inadimplencia(self):
        self._receber('1000.00')
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.status, 'parcial')
        self.assertEqual(self.receita.saldo_em_aberto, Decimal('1000.00'))
        self.assertFalse(self.receita.esta_quitada)
        self.assertIn(self.receita, receitas_inadimplentes_qs())

    def test_segundo_pagamento_soma_ao_primeiro(self):
        self._receber('1000.00')
        self._receber('1000.00')
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.valor_recebido, Decimal('2000.00'))
        self.assertEqual(self.receita.status, 'recebido')
        self.assertEqual(self.receita.saldo_em_aberto, Decimal('0.00'))
        self.assertNotIn(self.receita, receitas_inadimplentes_qs())

    def test_multa_juros_desconto_integram_o_saldo(self):
        self.receita.multa = Decimal('40.00')
        self.receita.juros = Decimal('10.00')
        self.receita.desconto = Decimal('50.00')
        self.receita.save()
        self.assertEqual(self.receita.valor_total_devido, Decimal('2000.00'))
        self._receber('2000.00')
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.status, 'recebido')
        # com multa maior, 2000 não quita
        self.receita.multa = Decimal('100.00')
        self.receita.save()
        self.assertEqual(self.receita.saldo_em_aberto, Decimal('60.00'))
        self.assertFalse(self.receita.esta_quitada)

    def test_cancelada_nao_aparece_como_divida(self):
        self.receita.status = 'cancelado'
        self.receita.save()
        self.assertEqual(self.receita.saldo_em_aberto, Decimal('0.00'))
        self.assertTrue(self.receita.esta_quitada)
        self.assertNotIn(self.receita, receitas_inadimplentes_qs())

    def test_recebimento_negativo_ou_acima_do_saldo_rejeitado(self):
        rec = self.RecebimentoReceita(
            receita=self.receita, data_recebimento=timezone.localdate(), valor=Decimal('-10.00'),
        )
        with self.assertRaises(ValidationError):
            rec.full_clean()
        rec2 = self.RecebimentoReceita(
            receita=self.receita, data_recebimento=timezone.localdate(), valor=Decimal('2000.01'),
        )
        with self.assertRaises(ValidationError):
            rec2.full_clean()

    def test_excluir_recebimento_reconsolida(self):
        r1 = self._receber('1000.00')
        self._receber('1000.00')
        r1.delete()
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.valor_recebido, Decimal('1000.00'))
        self.assertEqual(self.receita.status, 'parcial')

    def test_valor_recebido_legado_materializado_ao_receber(self):
        # dado legado: consolidado sem recebimentos
        self.receita.valor_recebido = Decimal('500.00')
        self.receita.save()
        self._receber('300.00')
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.valor_recebido, Decimal('800.00'))
        self.assertEqual(self.receita.recebimentos.count(), 2)

    def test_data_migration_recebimentos_iniciais_idempotente(self):
        import importlib
        from django.apps import apps as django_apps

        self.receita.valor_recebido = Decimal('700.00')
        self.receita.data_recebimento = self.ontem
        self.receita.save()

        mod = importlib.import_module('financeiro.migrations.0006_migrar_recebimentos_iniciais')
        mod.criar_recebimentos_iniciais(django_apps, None)
        mod.criar_recebimentos_iniciais(django_apps, None)  # segunda execução não duplica

        recebimentos = self.receita.recebimentos.all()
        self.assertEqual(recebimentos.count(), 1)
        self.assertEqual(recebimentos.first().valor, Decimal('700.00'))
        self.assertEqual(recebimentos.first().origem, 'migracao')


class OcupacaoRegrasTest(TestCase):
    """Regras defensivas de ocupação por status/encerramento (item 12)."""

    def setUp(self):
        self.imovel, self.locatario = _criar_base()
        self.hoje = timezone.localdate()

    def _ocupacao_atual(self):
        from financeiro.paineis import serie_ocupacao_mensal
        return serie_ocupacao_mensal(1, self.hoje)[0]['ocupados']

    def test_determinado_vigente_ocupa(self):
        _criar_contrato(self.imovel, self.locatario,
                        data_inicio=self.hoje - timedelta(days=100),
                        data_fim=self.hoje + timedelta(days=100))
        self.assertEqual(self._ocupacao_atual(), 1)

    def test_indeterminado_ativo_ocupa(self):
        _criar_contrato(self.imovel, self.locatario,
                        data_inicio=date(2020, 1, 1), data_fim=date(2020, 12, 31),
                        prazo_indeterminado=True)
        self.assertEqual(self._ocupacao_atual(), 1)

    def test_indeterminado_encerrado_sem_data_real_nao_ocupa_indefinidamente(self):
        _criar_contrato(self.imovel, self.locatario,
                        data_inicio=date(2020, 1, 1), data_fim=date(2020, 12, 31),
                        prazo_indeterminado=True, status='encerrado')
        self.assertEqual(self._ocupacao_atual(), 0)

    def test_rescindido_com_encerramento_no_meio_do_mes(self):
        from financeiro.paineis import serie_ocupacao_mensal
        meio_do_mes = self.hoje.replace(day=15)
        _criar_contrato(self.imovel, self.locatario,
                        data_inicio=date(2020, 1, 1), data_fim=date(2030, 12, 31),
                        status='rescindido', data_encerramento_real=meio_do_mes)
        # ocupa o mês do encerramento (parcial), mas não o seguinte
        self.assertEqual(serie_ocupacao_mensal(1, self.hoje)[0]['ocupados'], 1)
        proximo = (self.hoje.replace(day=1) + timedelta(days=32)).replace(day=1)
        self.assertEqual(serie_ocupacao_mensal(1, proximo)[0]['ocupados'], 0)

    def test_suspenso_nao_ocupa(self):
        _criar_contrato(self.imovel, self.locatario,
                        data_inicio=date(2020, 1, 1), data_fim=date(2030, 12, 31),
                        status='suspenso')
        self.assertEqual(self._ocupacao_atual(), 0)


# ─── Backup consistente e verificação (item 15) ───────────────────────────────

class BackupConsistenteTest(TestCase):
    def _sqlite_com_tabelas(self, pasta):
        import os
        import sqlite3
        db = os.path.join(pasta, 'real.sqlite3')
        con = sqlite3.connect(db)
        for tabela in ('patrimonio_imovel', 'patrimonio_contrato',
                       'financeiro_receitaaluguel', 'financeiro_despesa'):
            con.execute(f'CREATE TABLE {tabela} (id INTEGER PRIMARY KEY)')
        con.execute('INSERT INTO patrimonio_imovel (id) VALUES (1)')
        con.commit()
        con.close()
        return db

    def test_backup_sqlite_manifest_checksum_e_verificacao(self):
        import os
        import tempfile
        import zipfile
        from io import StringIO
        from django.core.management import call_command
        from django.test import override_settings

        pasta = tempfile.mkdtemp(prefix='bk_')
        db = self._sqlite_com_tabelas(pasta)
        media_vazia = tempfile.mkdtemp(prefix='bk_media_')

        databases = {'default': {'ENGINE': 'django.db.backends.sqlite3', 'NAME': db}}
        with override_settings(DATABASES=databases, MEDIA_ROOT=media_vazia):
            out = StringIO()
            call_command('backup_local', destino=pasta, stdout=out)

        zips = [f for f in os.listdir(pasta) if f.endswith('.zip')]
        self.assertEqual(len(zips), 1)
        zip_path = os.path.join(pasta, zips[0])
        # checksum do ZIP gravado ao lado
        self.assertTrue(os.path.exists(zip_path + '.sha256'))

        with zipfile.ZipFile(zip_path) as zf:
            self.assertIn('db.sqlite3', zf.namelist())
            manifest = zf.read('manifest.txt').decode('utf-8')
        self.assertIn('Engine do banco: django.db.backends.sqlite3', manifest)
        self.assertIn('SHA-256 do banco:', manifest)
        self.assertIn('Tamanho do banco (bytes):', manifest)
        # backup sem mídia funciona (aviso, não erro)
        self.assertIn('Arquivos de mídia: 0', manifest)

        out2 = StringIO()
        call_command('verificar_backup', zip_path, stdout=out2)
        saida = out2.getvalue()
        self.assertIn('Checksum SHA-256 do banco confere', saida)
        self.assertIn('patrimonio_imovel: 1 registro(s)', saida)
        self.assertIn('Backup válido', saida)

    def test_postgres_avisa_que_nao_fez_backup_do_banco(self):
        import tempfile
        from io import StringIO
        from django.core.management import call_command
        from django.test import override_settings

        databases = {'default': {'ENGINE': 'django.db.backends.postgresql', 'NAME': 'x'}}
        pasta = tempfile.mkdtemp(prefix='bk_pg_')
        with override_settings(DATABASES=databases, MEDIA_ROOT=tempfile.mkdtemp()):
            out, err = StringIO(), StringIO()
            call_command('backup_local', destino=pasta, stdout=out, stderr=err)
        self.assertIn('pg_dump', err.getvalue())
        self.assertIn('NÃO fez backup do banco', err.getvalue())


# ─── Rodada 3: reconsolidação de receitas canceladas (item 1) ─────────────────

class ReconsolidacaoCanceladaTest(TestCase):
    def setUp(self):
        self.imovel, self.locatario = _criar_base()
        self.contrato = _criar_contrato(
            self.imovel, self.locatario, data_inicio=date(2020, 1, 1), data_fim=date(2030, 12, 31),
        )
        self.receita = ReceitaAluguel.objects.create(
            contrato=self.contrato, imovel=self.imovel,
            competencia_mes=1, competencia_ano=2030,
            data_vencimento=date(2030, 1, 10), valor_previsto=Decimal('2000.00'), status='previsto',
        )

    def test_excluir_unico_recebimento_enquanto_cancelada_zera_consolidado(self):
        from financeiro.services import registrar_recebimento
        recebimento = registrar_recebimento(self.receita.pk, Decimal('500.00'), date(2030, 1, 15))
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.valor_recebido, Decimal('500.00'))

        self.receita.cancelar()
        recebimento.delete()

        self.receita.refresh_from_db()
        self.assertEqual(self.receita.status, 'cancelado')
        self.assertIsNone(self.receita.valor_recebido)
        self.assertIsNone(self.receita.data_recebimento)
        self.assertEqual(self.receita.recebimentos.count(), 0)

    def test_editar_valor_de_recebimento_enquanto_cancelada_reconsolida(self):
        from financeiro.services import registrar_recebimento
        recebimento = registrar_recebimento(self.receita.pk, Decimal('500.00'), date(2030, 1, 15))
        self.receita.cancelar()

        recebimento.valor = Decimal('450.00')
        recebimento.save()

        self.receita.refresh_from_db()
        self.assertEqual(self.receita.status, 'cancelado')
        self.assertEqual(self.receita.valor_recebido, Decimal('450.00'))

    def test_cancelar_excluir_recebimento_e_reabrir_nao_reaparece_fantasma(self):
        from financeiro.services import registrar_recebimento
        recebimento = registrar_recebimento(self.receita.pk, Decimal('500.00'), date(2030, 1, 15))
        self.receita.cancelar()
        recebimento.delete()

        self.receita.refresh_from_db()
        self.receita.reabrir()

        self.receita.refresh_from_db()
        self.assertIsNone(self.receita.valor_recebido)
        self.assertIn(self.receita.status, ('previsto', 'atrasado'))
        self.assertEqual(self.receita.saldo_em_aberto, Decimal('2000.00'))

    def test_cancelar_com_varios_recebimentos_excluir_um_e_reabrir(self):
        from financeiro.services import registrar_recebimento
        r1 = registrar_recebimento(self.receita.pk, Decimal('300.00'), date(2030, 1, 12))
        r2 = registrar_recebimento(self.receita.pk, Decimal('200.00'), date(2030, 1, 15))
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.valor_recebido, Decimal('500.00'))

        self.receita.cancelar()
        r2.delete()

        self.receita.refresh_from_db()
        self.assertEqual(self.receita.status, 'cancelado')
        self.assertEqual(self.receita.valor_recebido, Decimal('300.00'))
        self.assertEqual(self.receita.saldo_em_aberto, Decimal('0.00'))

        self.receita.reabrir()
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.status, 'parcial')
        self.assertEqual(self.receita.valor_recebido, Decimal('300.00'))
        self.assertEqual(self.receita.recebimentos.count(), 1)
        self.assertEqual(self.receita.recebimentos.get().pk, r1.pk)

    def test_nenhum_pagamento_legado_fantasma_apos_ciclo_cancelar_excluir_reabrir(self):
        """Garante que o valor consolidado nunca fica 'preso' após exclusão enquanto cancelada."""
        from financeiro.services import registrar_recebimento
        recebimento = registrar_recebimento(self.receita.pk, Decimal('2000.00'), date(2030, 1, 10))
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.status, 'recebido')

        self.receita.cancelar()
        recebimento.delete()
        self.receita.refresh_from_db()
        # a receita cancelada com o recebimento excluído não deve "lembrar" dos 2000
        self.assertIsNone(self.receita.valor_recebido)

        self.receita.reabrir()
        self.receita.refresh_from_db()
        self.assertNotEqual(self.receita.status, 'recebido')
        self.assertIsNone(self.receita.valor_recebido)
        self.assertEqual(self.receita.saldo_em_aberto, Decimal('2000.00'))


# ─── Rodada 3: service registrar_recebimento (item 2) ──────────────────────────

class RegistrarRecebimentoServiceTest(TestCase):
    def setUp(self):
        self.imovel, self.locatario = _criar_base()
        self.contrato = _criar_contrato(
            self.imovel, self.locatario, data_inicio=date(2020, 1, 1), data_fim=date(2030, 12, 31),
        )
        self.receita = ReceitaAluguel.objects.create(
            contrato=self.contrato, imovel=self.imovel,
            competencia_mes=2, competencia_ano=2030,
            data_vencimento=date(2030, 2, 10), valor_previsto=Decimal('1000.00'), status='previsto',
        )

    def test_receita_cancelada_rejeitada(self):
        from financeiro.services import registrar_recebimento
        self.receita.cancelar()
        with self.assertRaises(ValidationError):
            registrar_recebimento(self.receita.pk, Decimal('100.00'), date(2030, 2, 15))

    def test_valor_acima_do_saldo_rejeitado(self):
        from financeiro.services import registrar_recebimento
        with self.assertRaises(ValidationError):
            registrar_recebimento(self.receita.pk, Decimal('1000.01'), date(2030, 2, 15))
        self.receita.refresh_from_db()
        self.assertIsNone(self.receita.valor_recebido)

    def test_valor_zero_ou_negativo_rejeitado(self):
        from financeiro.services import registrar_recebimento
        with self.assertRaises(ValidationError):
            registrar_recebimento(self.receita.pk, Decimal('0'), date(2030, 2, 15))
        with self.assertRaises(ValidationError):
            registrar_recebimento(self.receita.pk, Decimal('-10.00'), date(2030, 2, 15))

    def test_dois_pagamentos_parciais_validos_e_quitacao_do_restante(self):
        from financeiro.services import registrar_recebimento
        registrar_recebimento(self.receita.pk, Decimal('400.00'), date(2030, 2, 12))
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.status, 'parcial')
        self.assertEqual(self.receita.saldo_em_aberto, Decimal('600.00'))

        registrar_recebimento(self.receita.pk, Decimal('600.00'), date(2030, 2, 15))
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.status, 'recebido')
        self.assertEqual(self.receita.saldo_em_aberto, Decimal('0.00'))
        self.assertEqual(self.receita.recebimentos.count(), 2)

    def test_rollback_integral_em_caso_de_erro(self):
        """Se o full_clean() rejeita, nenhum RecebimentoReceita é criado."""
        from financeiro.services import registrar_recebimento
        with self.assertRaises(ValidationError):
            registrar_recebimento(self.receita.pk, Decimal('99999.00'), date(2030, 2, 15))
        self.assertEqual(self.receita.recebimentos.count(), 0)
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.status, 'previsto')

    def test_retorna_o_recebimento_criado(self):
        from financeiro.services import registrar_recebimento
        recebimento = registrar_recebimento(
            self.receita.pk, Decimal('250.00'), date(2030, 2, 12),
            observacoes='teste', origem='manual',
        )
        self.assertEqual(recebimento.pk, self.receita.recebimentos.get().pk)
        self.assertEqual(recebimento.valor, Decimal('250.00'))
        self.assertEqual(recebimento.observacoes, 'teste')


class AtualizarRecebimentosDaReceitaTest(TestCase):
    """
    Item 1 (rodada pós-revisão): atualizar_recebimentos_da_receita() trata o
    formset inteiro do inline do Admin como uma única operação financeira.
    """
    def setUp(self):
        self.imovel, self.locatario = _criar_base()
        self.contrato = _criar_contrato(
            self.imovel, self.locatario, data_inicio=date(2020, 1, 1), data_fim=date(2030, 12, 31),
        )
        self.receita = ReceitaAluguel.objects.create(
            contrato=self.contrato, imovel=self.imovel,
            competencia_mes=2, competencia_ano=2030,
            data_vencimento=date(2030, 2, 10), valor_previsto=Decimal('1000.00'), status='previsto',
        )

    def test_duas_novas_linhas_que_somadas_excedem_saldo_sao_rejeitadas(self):
        from financeiro.services import atualizar_recebimentos_da_receita
        novos = [
            {'data_recebimento': date(2030, 2, 12), 'valor': Decimal('600.00'), 'observacoes': ''},
            {'data_recebimento': date(2030, 2, 13), 'valor': Decimal('600.00'), 'observacoes': ''},
        ]
        with self.assertRaises(ValidationError):
            atualizar_recebimentos_da_receita(self.receita.pk, novos=novos)
        self.assertEqual(self.receita.recebimentos.count(), 0)
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.status, 'previsto')

    def test_uma_linha_de_600_e_outra_de_400_sao_aceitas(self):
        from financeiro.services import atualizar_recebimentos_da_receita
        novos = [
            {'data_recebimento': date(2030, 2, 12), 'valor': Decimal('600.00'), 'observacoes': ''},
            {'data_recebimento': date(2030, 2, 13), 'valor': Decimal('400.00'), 'observacoes': 'quitação'},
        ]
        atualizar_recebimentos_da_receita(self.receita.pk, novos=novos)
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.recebimentos.count(), 2)
        self.assertEqual(self.receita.valor_recebido, Decimal('1000.00'))
        self.assertEqual(self.receita.status, 'recebido')

    def test_editar_dois_recebimentos_simultaneamente_respeita_total_final(self):
        from financeiro.services import atualizar_recebimentos_da_receita
        r1 = RecebimentoReceita.objects.create(
            receita=self.receita, data_recebimento=date(2030, 2, 5), valor=Decimal('500.00'), origem='manual',
        )
        r2 = RecebimentoReceita.objects.create(
            receita=self.receita, data_recebimento=date(2030, 2, 6), valor=Decimal('500.00'), origem='manual',
        )
        # 300 + 700 = 1000: válido em conjunto, mas r2 isolado (700) excederia
        # o saldo se validado individualmente contra o saldo atual (0 em aberto).
        alterados = [
            {'pk': r1.pk, 'data_recebimento': r1.data_recebimento, 'valor': Decimal('300.00'), 'observacoes': ''},
            {'pk': r2.pk, 'data_recebimento': r2.data_recebimento, 'valor': Decimal('700.00'), 'observacoes': ''},
        ]
        atualizar_recebimentos_da_receita(self.receita.pk, alterados=alterados)
        r1.refresh_from_db(); r2.refresh_from_db()
        self.assertEqual(r1.valor, Decimal('300.00'))
        self.assertEqual(r2.valor, Decimal('700.00'))
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.valor_recebido, Decimal('1000.00'))
        self.assertEqual(self.receita.status, 'recebido')

    def test_excluir_um_recebimento_e_aumentar_outro_usa_saldo_final_correto(self):
        from financeiro.services import atualizar_recebimentos_da_receita
        r1 = RecebimentoReceita.objects.create(
            receita=self.receita, data_recebimento=date(2030, 2, 5), valor=Decimal('300.00'), origem='manual',
        )
        r2 = RecebimentoReceita.objects.create(
            receita=self.receita, data_recebimento=date(2030, 2, 6), valor=Decimal('300.00'), origem='manual',
        )
        # Exclui r1 (300) e aumenta r2 para 1000 — total final = 1000, válido.
        atualizar_recebimentos_da_receita(
            self.receita.pk,
            alterados=[{'pk': r2.pk, 'data_recebimento': r2.data_recebimento, 'valor': Decimal('1000.00'), 'observacoes': ''}],
            excluidos=[r1.pk],
        )
        self.assertFalse(RecebimentoReceita.objects.filter(pk=r1.pk).exists())
        r2.refresh_from_db()
        self.assertEqual(r2.valor, Decimal('1000.00'))
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.valor_recebido, Decimal('1000.00'))
        self.assertEqual(self.receita.status, 'recebido')

    def test_erro_em_uma_linha_causa_rollback_de_todas(self):
        from financeiro.services import atualizar_recebimentos_da_receita
        r1 = RecebimentoReceita.objects.create(
            receita=self.receita, data_recebimento=date(2030, 2, 5), valor=Decimal('200.00'), origem='manual',
        )
        novos = [{'data_recebimento': date(2030, 2, 12), 'valor': Decimal('300.00'), 'observacoes': ''}]
        alterados = [{'pk': r1.pk, 'data_recebimento': r1.data_recebimento, 'valor': Decimal('-50.00'), 'observacoes': ''}]
        with self.assertRaises(ValidationError):
            atualizar_recebimentos_da_receita(self.receita.pk, novos=novos, alterados=alterados)
        # nada foi persistido: nem a linha nova, nem a alteração da existente
        self.assertEqual(self.receita.recebimentos.count(), 1)
        r1.refresh_from_db()
        self.assertEqual(r1.valor, Decimal('200.00'))

    def test_receita_cancelada_rejeita_inclusao(self):
        from financeiro.services import atualizar_recebimentos_da_receita
        self.receita.cancelar()
        novos = [{'data_recebimento': date(2030, 2, 12), 'valor': Decimal('100.00'), 'observacoes': ''}]
        with self.assertRaises(ValidationError):
            atualizar_recebimentos_da_receita(self.receita.pk, novos=novos)
        self.assertEqual(self.receita.recebimentos.count(), 0)

    def test_recebimento_de_conciliacao_nao_pode_ser_editado(self):
        from financeiro.services import atualizar_recebimentos_da_receita
        r1 = RecebimentoReceita.objects.create(
            receita=self.receita, data_recebimento=date(2030, 2, 5), valor=Decimal('500.00'), origem='conciliacao',
        )
        alterados = [{'pk': r1.pk, 'data_recebimento': r1.data_recebimento, 'valor': Decimal('600.00'), 'observacoes': ''}]
        with self.assertRaises(ValidationError) as ctx:
            atualizar_recebimentos_da_receita(self.receita.pk, alterados=alterados)
        self.assertIn('conciliação', str(ctx.exception))
        r1.refresh_from_db()
        self.assertEqual(r1.valor, Decimal('500.00'))

    def test_recebimento_de_conciliacao_nao_pode_ser_excluido(self):
        from financeiro.services import atualizar_recebimentos_da_receita
        r1 = RecebimentoReceita.objects.create(
            receita=self.receita, data_recebimento=date(2030, 2, 5), valor=Decimal('500.00'), origem='conciliacao',
        )
        with self.assertRaises(ValidationError) as ctx:
            atualizar_recebimentos_da_receita(self.receita.pk, excluidos=[r1.pk])
        self.assertIn('conciliação', str(ctx.exception))
        self.assertTrue(RecebimentoReceita.objects.filter(pk=r1.pk).exists())

    def test_recebimentos_manuais_continuam_editaveis(self):
        from financeiro.services import atualizar_recebimentos_da_receita
        r1 = RecebimentoReceita.objects.create(
            receita=self.receita, data_recebimento=date(2030, 2, 5), valor=Decimal('500.00'), origem='manual',
        )
        alterados = [{'pk': r1.pk, 'data_recebimento': date(2030, 2, 6), 'valor': Decimal('450.00'), 'observacoes': 'ajuste'}]
        atualizar_recebimentos_da_receita(self.receita.pk, alterados=alterados)
        r1.refresh_from_db()
        self.assertEqual(r1.valor, Decimal('450.00'))
        self.assertEqual(r1.data_recebimento, date(2030, 2, 6))
        self.assertEqual(r1.observacoes, 'ajuste')

    def test_consolidado_final_corresponde_exatamente_a_soma_das_linhas(self):
        from financeiro.services import atualizar_recebimentos_da_receita
        r1 = RecebimentoReceita.objects.create(
            receita=self.receita, data_recebimento=date(2030, 2, 5), valor=Decimal('300.00'), origem='manual',
        )
        novos = [{'data_recebimento': date(2030, 2, 12), 'valor': Decimal('250.00'), 'observacoes': ''}]
        alterados = [{'pk': r1.pk, 'data_recebimento': r1.data_recebimento, 'valor': Decimal('350.00'), 'observacoes': ''}]
        atualizar_recebimentos_da_receita(self.receita.pk, novos=novos, alterados=alterados)
        self.receita.refresh_from_db()
        soma_persistida = sum((r.valor for r in self.receita.recebimentos.all()), Decimal('0.00'))
        self.assertEqual(soma_persistida, Decimal('600.00'))
        self.assertEqual(self.receita.valor_recebido, soma_persistida)


class RecebimentoInlineAdminFormsetTest(TestCase):
    """
    Item 1: exercita o formset REAL do inline (não uma versão simplificada),
    para provar que os campos disabled=True de fato bloqueiam alteração e
    exclusão de recebimentos de conciliação no nível do Django forms, e que
    ReceitaAluguelAdmin.save_formset delega ao service corretamente.
    """
    def setUp(self):
        from django.contrib.auth.models import User
        self.superuser = User.objects.create_superuser('admin_receb', 'a@a.com', 'pass')
        self.imovel, self.locatario = _criar_base()
        self.contrato = _criar_contrato(
            self.imovel, self.locatario, data_inicio=date(2020, 1, 1), data_fim=date(2030, 12, 31),
        )
        self.receita = ReceitaAluguel.objects.create(
            contrato=self.contrato, imovel=self.imovel,
            competencia_mes=2, competencia_ano=2030,
            data_vencimento=date(2030, 2, 10), valor_previsto=Decimal('1000.00'), status='previsto',
        )

    def _request(self):
        from django.contrib.messages.storage.fallback import FallbackStorage
        from django.test import RequestFactory
        request = RequestFactory().post('/admin/financeiro/receitaaluguel/%d/change/' % self.receita.pk)
        request.user = self.superuser
        request.session = {}
        request._messages = FallbackStorage(request)
        return request

    def _formset_class(self, request):
        from django.contrib import admin as django_admin
        from financeiro.admin import RecebimentoReceitaInline
        inline = RecebimentoReceitaInline(ReceitaAluguel, django_admin.site)
        return inline.get_formset(request, obj=self.receita)

    def _post_data(self, prefix, linhas):
        inicial = sum(1 for linha in linhas if linha.get('id'))
        dados = {
            f'{prefix}-TOTAL_FORMS': str(len(linhas)),
            f'{prefix}-INITIAL_FORMS': str(inicial),
            f'{prefix}-MIN_NUM_FORMS': '0',
            f'{prefix}-MAX_NUM_FORMS': '1000',
        }
        for i, linha in enumerate(linhas):
            dados[f'{prefix}-{i}-id'] = str(linha.get('id', ''))
            dados[f'{prefix}-{i}-data_recebimento'] = linha.get('data_recebimento', '')
            dados[f'{prefix}-{i}-valor'] = linha.get('valor', '')
            dados[f'{prefix}-{i}-observacoes'] = linha.get('observacoes', '')
            if linha.get('DELETE'):
                dados[f'{prefix}-{i}-DELETE'] = 'on'
        return dados

    def test_recebimento_de_conciliacao_ignora_alteracao_no_nivel_do_formset(self):
        r1 = RecebimentoReceita.objects.create(
            receita=self.receita, data_recebimento=date(2030, 2, 5), valor=Decimal('500.00'), origem='conciliacao',
        )
        request = self._request()
        FormSetClass = self._formset_class(request)
        prefix = FormSetClass.get_default_prefix()
        dados = self._post_data(prefix, [
            {'id': r1.pk, 'data_recebimento': '2030-02-09', 'valor': '999.00', 'observacoes': 'tentativa'},
        ])
        formset = FormSetClass(data=dados, instance=self.receita, prefix=prefix)
        self.assertTrue(formset.is_valid(), formset.errors)
        formset.save(commit=False)
        # campo disabled: o valor submetido é ignorado — nenhuma mudança detectada
        self.assertEqual(formset.changed_objects, [])

    def test_recebimento_de_conciliacao_ignora_exclusao_no_nivel_do_formset(self):
        r1 = RecebimentoReceita.objects.create(
            receita=self.receita, data_recebimento=date(2030, 2, 5), valor=Decimal('500.00'), origem='conciliacao',
        )
        request = self._request()
        FormSetClass = self._formset_class(request)
        prefix = FormSetClass.get_default_prefix()
        dados = self._post_data(prefix, [
            {'id': r1.pk, 'data_recebimento': '2030-02-05', 'valor': '500.00', 'observacoes': '', 'DELETE': True},
        ])
        formset = FormSetClass(data=dados, instance=self.receita, prefix=prefix)
        self.assertTrue(formset.is_valid(), formset.errors)
        formset.save(commit=False)
        # checkbox DELETE disabled: exclusão nunca é considerada
        self.assertEqual(formset.deleted_objects, [])

    def test_formset_duas_novas_linhas_excedem_saldo_fica_invalido_sem_persistir(self):
        """
        Item 1: a inconsistência do lote agora invalida o FORMSET no ciclo de
        validação do Django (formset.is_valid() == False) — antes de
        qualquer persistência, e não mais só depois via save_formset()
        capturando ValidationError e emitindo messages.error. No Admin real,
        um formset inválido nunca chega a save_model()/save_formset().
        """
        request = self._request()
        FormSetClass = self._formset_class(request)
        prefix = FormSetClass.get_default_prefix()
        dados = self._post_data(prefix, [
            {'id': '', 'data_recebimento': '2030-02-12', 'valor': '600.00', 'observacoes': ''},
            {'id': '', 'data_recebimento': '2030-02-13', 'valor': '600.00', 'observacoes': ''},
        ])
        formset = FormSetClass(data=dados, instance=self.receita, prefix=prefix)
        self.assertFalse(formset.is_valid())
        self.assertTrue(any('excederia' in erro for erro in formset.non_form_errors()))
        self.assertEqual(self.receita.recebimentos.count(), 0)

    def test_save_formset_linhas_validas_persiste(self):
        from types import SimpleNamespace
        from django.contrib import admin as django_admin
        from financeiro.admin import ReceitaAluguelAdmin

        request = self._request()
        FormSetClass = self._formset_class(request)
        prefix = FormSetClass.get_default_prefix()
        dados = self._post_data(prefix, [
            {'id': '', 'data_recebimento': '2030-02-12', 'valor': '600.00', 'observacoes': ''},
            {'id': '', 'data_recebimento': '2030-02-13', 'valor': '400.00', 'observacoes': ''},
        ])
        formset = FormSetClass(data=dados, instance=self.receita, prefix=prefix)
        self.assertTrue(formset.is_valid(), formset.errors)

        admin_instance = ReceitaAluguelAdmin(ReceitaAluguel, django_admin.site)
        form = SimpleNamespace(instance=self.receita)
        admin_instance.save_formset(request, form, formset, change=True)

        self.receita.refresh_from_db()
        self.assertEqual(self.receita.recebimentos.count(), 2)
        self.assertEqual(self.receita.valor_recebido, Decimal('1000.00'))
        self.assertEqual(self.receita.status, 'recebido')


class RecebimentoInlineAdminIntegracaoTest(TestCase):
    """
    Item 1 (rodada final): testes de integração com SUBMISSÃO REAL ao
    change form do Django Admin via Client.post() — não chamada direta a
    save_formset(). Prova que o lote inválido invalida o formset no ciclo
    correto (antes de qualquer save_model()/save_related()): nada é
    persistido (nem multa do form principal, nem recebimentos), a página é
    reapresentada com os valores submetidos e sem mensagem de sucesso.
    """
    def setUp(self):
        from django.contrib.auth.models import User
        self.superuser = User.objects.create_superuser('admin_receb_full', 'a@a.com', 'pass')
        self.client = Client()
        self.client.login(username='admin_receb_full', password='pass')
        self.imovel, self.locatario = _criar_base()
        self.contrato = _criar_contrato(
            self.imovel, self.locatario, data_inicio=date(2020, 1, 1), data_fim=date(2030, 12, 31),
        )
        self.receita = ReceitaAluguel.objects.create(
            contrato=self.contrato, imovel=self.imovel,
            competencia_mes=2, competencia_ano=2030,
            data_vencimento=date(2030, 2, 10), valor_previsto=Decimal('1000.00'), status='previsto',
        )
        self.url = reverse('admin:financeiro_receitaaluguel_change', args=[self.receita.pk])

    def _get_form_e_formsets(self):
        resposta = self.client.get(self.url)
        return resposta, resposta.context['adminform'], resposta.context['inline_admin_formsets']

    def _dados_base(self, adminform, inline_formsets):
        """Monta o POST a partir do changeform REAL renderizado — nunca
        adivinha nomes de campo/prefixo (evita testes frágeis). Inclui os
        dados de CADA linha existente de cada inline (id/valor/etc.), não só
        os totais do management form — necessário para reenviar linhas
        existentes inalteradas junto com as que o teste efetivamente altera."""
        dados = {}
        for name, field in adminform.form.fields.items():
            valor = adminform.form.initial.get(name, field.initial)
            if valor is None:
                valor = ''
            if hasattr(valor, 'pk'):
                valor = valor.pk
            dados[name] = valor
        for iaf in inline_formsets:
            fs = iaf.formset
            mgmt = fs.management_form.initial
            for campo in ('TOTAL_FORMS', 'INITIAL_FORMS', 'MIN_NUM_FORMS', 'MAX_NUM_FORMS'):
                dados[f'{fs.prefix}-{campo}'] = mgmt.get(campo, 0)
            for form in fs.forms:
                for name, field in form.fields.items():
                    valor = form.initial.get(name, field.initial)
                    if valor is None:
                        valor = ''
                    if hasattr(valor, 'pk'):
                        valor = valor.pk
                    dados[form.add_prefix(name)] = valor
        return dados

    def _formset_recebimentos(self, inline_formsets):
        for iaf in inline_formsets:
            if iaf.formset.model is RecebimentoReceita:
                return iaf.formset
        raise AssertionError('inline de recebimentos não encontrado no changeform')

    def _adicionar_novo_recebimento(self, dados, prefix, indice, valor, data='2030-02-12', obs=''):
        dados[f'{prefix}-{indice}-id'] = ''
        dados[f'{prefix}-{indice}-data_recebimento'] = data
        dados[f'{prefix}-{indice}-valor'] = valor
        dados[f'{prefix}-{indice}-observacoes'] = obs

    def test_alterar_multa_e_duas_linhas_de_600_reapresenta_form_com_erro(self):
        """Cenário 1 (10 testes obrigatórios do item 1, combinados aqui)."""
        _resposta, adminform, inline_formsets = self._get_form_e_formsets()
        fs = self._formset_recebimentos(inline_formsets)
        dados = self._dados_base(adminform, inline_formsets)
        dados['multa'] = '50.00'
        dados[f'{fs.prefix}-TOTAL_FORMS'] = '2'
        self._adicionar_novo_recebimento(dados, fs.prefix, 0, '600.00', '2030-02-12')
        self._adicionar_novo_recebimento(dados, fs.prefix, 1, '600.00', '2030-02-13')

        resposta = self.client.post(self.url, dados)

        # reapresenta o formulário com erro (200), não redireciona (302)
        self.assertEqual(resposta.status_code, 200)
        self.receita.refresh_from_db()
        # multa não persistida
        self.assertEqual(self.receita.multa, Decimal('0.00'))
        # nenhum recebimento criado
        self.assertEqual(self.receita.recebimentos.count(), 0)
        # sem mensagem de sucesso do Admin
        conteudo = resposta.content.decode()
        self.assertNotIn('foi alterad', conteudo.lower())
        self.assertNotIn('was changed successfully', conteudo.lower())
        # valores submetidos continuam visíveis no formulário reapresentado
        self.assertIn('600.00', conteudo)
        self.assertIn('50.00', conteudo)
        # erro do lote aparece no formset de recebimentos
        self.assertTrue(any('excederia' in erro for erro in resposta.context['inline_admin_formsets'][1].formset.non_form_errors()))

    def test_seiscentos_mais_quatrocentos_salva_tudo(self):
        """Cenário 9: R$ 600 + R$ 400 (saldo exato de R$ 1.000) deve quitar a receita."""
        _resposta, adminform, inline_formsets = self._get_form_e_formsets()
        fs = self._formset_recebimentos(inline_formsets)
        dados = self._dados_base(adminform, inline_formsets)
        dados[f'{fs.prefix}-TOTAL_FORMS'] = '2'
        self._adicionar_novo_recebimento(dados, fs.prefix, 0, '600.00', '2030-02-12')
        self._adicionar_novo_recebimento(dados, fs.prefix, 1, '400.00', '2030-02-13')

        resposta = self.client.post(self.url, dados)

        self.assertEqual(resposta.status_code, 302)  # redireciona = sucesso
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.recebimentos.count(), 2)
        self.assertEqual(self.receita.valor_recebido, Decimal('1000.00'))
        self.assertEqual(self.receita.status, 'recebido')

    def test_alteracao_e_exclusao_simultaneas_validadas_pelo_estado_final(self):
        """Cenário 10: dois recebimentos existentes (R$300+R$300, saldo livre
        R$400) — a submissão aumenta um para R$650 e exclui o outro; o
        estado FINAL (só R$650) cabe no saldo, mesmo que a soma intermediária
        (existente 300 + alterado 650) pareça exceder olhando linha a linha."""
        from financeiro.services import atualizar_recebimentos_da_receita
        r1 = RecebimentoReceita.objects.create(
            receita=self.receita, data_recebimento=date(2030, 2, 5), valor=Decimal('300.00'), origem='manual',
        )
        r2 = RecebimentoReceita.objects.create(
            receita=self.receita, data_recebimento=date(2030, 2, 6), valor=Decimal('300.00'), origem='manual',
        )
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.valor_recebido, Decimal('600.00'))

        _resposta, adminform, inline_formsets = self._get_form_e_formsets()
        fs = self._formset_recebimentos(inline_formsets)
        dados = self._dados_base(adminform, inline_formsets)
        # dois formulários existentes (r1, r2) já vêm no INITIAL_FORMS/TOTAL_FORMS
        indice_r1 = next(i for i, f in enumerate(fs.forms) if f.instance.pk == r1.pk)
        indice_r2 = next(i for i, f in enumerate(fs.forms) if f.instance.pk == r2.pk)
        dados[f'{fs.prefix}-{indice_r1}-valor'] = '650.00'
        dados[f'{fs.prefix}-{indice_r2}-DELETE'] = 'on'

        resposta = self.client.post(self.url, dados)

        self.assertEqual(resposta.status_code, 302)
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.recebimentos.count(), 1)
        r1.refresh_from_db()
        self.assertEqual(r1.valor, Decimal('650.00'))
        self.assertEqual(self.receita.valor_recebido, Decimal('650.00'))


class HistoricoRecebimentosEmLoteTest(TestCase):
    """
    Item 2 (rodada final): novos recebimentos do lote são salvos com save()
    individual (não bulk_create()) — o histórico (django-simple-history) é
    preservado por linha, mesmo que a reconsolidação da receita só aconteça
    uma vez ao final do lote.
    """
    def setUp(self):
        from financeiro.models import HistoricalRecebimentoReceita
        self.HistoricalRecebimentoReceita = HistoricalRecebimentoReceita
        self.usuario = User.objects.create_user('lote_hist', password='pass')
        self.imovel, self.locatario = _criar_base()
        self.contrato = _criar_contrato(
            self.imovel, self.locatario, data_inicio=date(2020, 1, 1), data_fim=date(2030, 12, 31),
        )
        self.receita = ReceitaAluguel.objects.create(
            contrato=self.contrato, imovel=self.imovel,
            competencia_mes=2, competencia_ano=2030,
            data_vencimento=date(2030, 2, 10), valor_previsto=Decimal('1000.00'), status='previsto',
        )

    def test_novo_recebimento_possui_historico_de_criacao(self):
        from financeiro.services import atualizar_recebimentos_da_receita
        novos = [{'data_recebimento': date(2030, 2, 12), 'valor': Decimal('300.00'), 'observacoes': ''}]
        atualizar_recebimentos_da_receita(self.receita.pk, novos=novos, usuario=self.usuario)
        rec = self.receita.recebimentos.get()
        historico = self.HistoricalRecebimentoReceita.objects.filter(id=rec.pk, history_type='+')
        self.assertEqual(historico.count(), 1)
        self.assertEqual(historico.first().criado_por_id, self.usuario.pk)

    def test_dois_recebimentos_na_mesma_submissao_possuem_historicos_separados(self):
        from financeiro.services import atualizar_recebimentos_da_receita
        novos = [
            {'data_recebimento': date(2030, 2, 12), 'valor': Decimal('300.00'), 'observacoes': ''},
            {'data_recebimento': date(2030, 2, 13), 'valor': Decimal('300.00'), 'observacoes': ''},
        ]
        atualizar_recebimentos_da_receita(self.receita.pk, novos=novos)
        pks = list(self.receita.recebimentos.values_list('pk', flat=True))
        self.assertEqual(len(pks), 2)
        for pk in pks:
            self.assertEqual(
                self.HistoricalRecebimentoReceita.objects.filter(id=pk, history_type='+').count(), 1,
            )

    def test_alteracao_gera_historico(self):
        from financeiro.services import atualizar_recebimentos_da_receita
        r1 = RecebimentoReceita.objects.create(
            receita=self.receita, data_recebimento=date(2030, 2, 5), valor=Decimal('300.00'), origem='manual',
        )
        alterados = [{'pk': r1.pk, 'data_recebimento': r1.data_recebimento, 'valor': Decimal('350.00'), 'observacoes': ''}]
        atualizar_recebimentos_da_receita(self.receita.pk, alterados=alterados)
        self.assertEqual(
            self.HistoricalRecebimentoReceita.objects.filter(id=r1.pk, history_type='~').count(), 1,
        )

    def test_exclusao_gera_historico(self):
        from financeiro.services import atualizar_recebimentos_da_receita
        r1 = RecebimentoReceita.objects.create(
            receita=self.receita, data_recebimento=date(2030, 2, 5), valor=Decimal('300.00'), origem='manual',
        )
        pk = r1.pk
        atualizar_recebimentos_da_receita(self.receita.pk, excluidos=[pk])
        self.assertFalse(RecebimentoReceita.objects.filter(pk=pk).exists())
        self.assertEqual(
            self.HistoricalRecebimentoReceita.objects.filter(id=pk, history_type='-').count(), 1,
        )

    def test_lote_rejeitado_nao_gera_historico(self):
        from financeiro.services import atualizar_recebimentos_da_receita
        novos = [
            {'data_recebimento': date(2030, 2, 12), 'valor': Decimal('600.00'), 'observacoes': ''},
            {'data_recebimento': date(2030, 2, 13), 'valor': Decimal('600.00'), 'observacoes': ''},
        ]
        with self.assertRaises(ValidationError):
            atualizar_recebimentos_da_receita(self.receita.pk, novos=novos)
        self.assertEqual(RecebimentoReceita.objects.filter(receita=self.receita).count(), 0)
        self.assertEqual(
            self.HistoricalRecebimentoReceita.objects.filter(receita_id=self.receita.pk).count(), 0,
        )

    def test_rollback_nao_deixa_registros_historicos_orfaos(self):
        """
        Simula uma falha REAL no meio da persistência (não uma rejeição de
        validação, que nunca chega a salvar nada) — a primeira linha chega a
        ser salva (e gera histórico de criação) antes da segunda falhar; o
        transaction.atomic() do service reverte TUDO, inclusive o histórico
        da primeira linha, já que o sinal post_save do simple-history grava
        dentro da MESMA transação.
        """
        from unittest.mock import patch
        from financeiro.services import atualizar_recebimentos_da_receita

        original_save = RecebimentoReceita.save
        chamadas = {'n': 0}

        def save_com_falha_na_segunda(self, *args, **kwargs):
            chamadas['n'] += 1
            if chamadas['n'] == 2:
                raise RuntimeError('falha simulada no meio do lote')
            return original_save(self, *args, **kwargs)

        novos = [
            {'data_recebimento': date(2030, 2, 12), 'valor': Decimal('300.00'), 'observacoes': ''},
            {'data_recebimento': date(2030, 2, 13), 'valor': Decimal('300.00'), 'observacoes': ''},
        ]
        with patch.object(RecebimentoReceita, 'save', save_com_falha_na_segunda):
            with self.assertRaises(RuntimeError):
                atualizar_recebimentos_da_receita(self.receita.pk, novos=novos)

        self.assertEqual(RecebimentoReceita.objects.filter(receita=self.receita).count(), 0)
        self.assertEqual(
            self.HistoricalRecebimentoReceita.objects.filter(receita_id=self.receita.pk).count(), 0,
        )

    def test_receita_reconsolidada_uma_unica_vez_por_lote(self):
        from unittest.mock import patch
        from financeiro.services import atualizar_recebimentos_da_receita

        original = ReceitaAluguel.recalcular_recebimentos
        chamadas = []

        def contador(self, *args, **kwargs):
            chamadas.append(self.pk)
            return original(self, *args, **kwargs)

        novos = [
            {'data_recebimento': date(2030, 2, 12), 'valor': Decimal('300.00'), 'observacoes': ''},
            {'data_recebimento': date(2030, 2, 13), 'valor': Decimal('300.00'), 'observacoes': ''},
        ]
        with patch.object(ReceitaAluguel, 'recalcular_recebimentos', contador):
            atualizar_recebimentos_da_receita(self.receita.pk, novos=novos)

        self.assertEqual(chamadas.count(self.receita.pk), 1)
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.valor_recebido, Decimal('600.00'))


class RegistrarRecebimentoConcorrenciaTest(TransactionTestCase):
    """
    Concorrência real (select_for_update bloqueando de fato) só é garantida
    pelo backend no PostgreSQL — o SQLite não suporta locking de linha
    (connection.features.has_select_for_update é False) e apenas ignora a
    cláusula, então o teste é pulado nesse backend (skipUnlessDBFeature é o
    padrão do próprio Django para este cenário). Roda de verdade no job
    Postgres do CI.
    """

    def setUp(self):
        self.imovel, self.locatario = _criar_base()
        self.contrato = _criar_contrato(
            self.imovel, self.locatario, data_inicio=date(2020, 1, 1), data_fim=date(2030, 12, 31),
        )
        self.receita = ReceitaAluguel.objects.create(
            contrato=self.contrato, imovel=self.imovel,
            competencia_mes=3, competencia_ano=2030,
            data_vencimento=date(2030, 3, 10), valor_previsto=Decimal('1000.00'), status='previsto',
        )

    @skipUnlessDBFeature('has_select_for_update')
    def test_duas_tentativas_simultaneas_nao_estouram_o_saldo(self):
        import threading
        from django.db import connections
        from financeiro.services import registrar_recebimento

        resultados = []
        barreira = threading.Barrier(2)

        def tentar():
            barreira.wait()
            try:
                registrar_recebimento(self.receita.pk, Decimal('700.00'), date(2030, 3, 15))
                resultados.append('ok')
            except ValidationError:
                resultados.append('rejeitado')
            finally:
                connections.close_all()

        t1 = threading.Thread(target=tentar)
        t2 = threading.Thread(target=tentar)
        t1.start()
        t2.start()
        t1.join()
        t2.join()

        self.assertEqual(sorted(resultados), ['ok', 'rejeitado'])
        self.receita.refresh_from_db()
        self.assertEqual(self.receita.valor_recebido, Decimal('700.00'))
        self.assertEqual(self.receita.recebimentos.count(), 1)
        total = self.receita.recebimentos.aggregate(models.Sum('valor'))['valor__sum']
        self.assertLessEqual(total, Decimal('1000.00'))


# ─── Rodada 3 (item 4): caixa de receita cancelada continua contando ───────────

class CaixaReceitaCanceladaTest(TestCase):
    """
    Cenário do item 4: receita de R$ 2.000, recebimento de R$ 500,
    cancelamento do saldo restante — caixa e relatórios devem continuar
    mostrando R$ 500; saldo em aberto deve ser zero; receita não deve
    aparecer como inadimplente.
    """

    def setUp(self):
        self.imovel, self.locatario = _criar_base()
        self.contrato = _criar_contrato(
            self.imovel, self.locatario, data_inicio=date(2020, 1, 1), data_fim=date(2030, 12, 31),
        )
        self.hoje = timezone.localdate()
        self.receita = ReceitaAluguel.objects.create(
            contrato=self.contrato, imovel=self.imovel,
            competencia_mes=self.hoje.month, competencia_ano=self.hoje.year,
            data_vencimento=self.hoje - timedelta(days=5), valor_previsto=Decimal('2000.00'), status='previsto',
        )
        from financeiro.services import registrar_recebimento
        registrar_recebimento(self.receita.pk, Decimal('500.00'), self.hoje)
        self.receita.cancelar()
        self.receita.refresh_from_db()

    def test_saldo_zero_e_nao_e_inadimplente(self):
        self.assertEqual(self.receita.saldo_em_aberto, Decimal('0.00'))
        self.assertNotIn(self.receita, receitas_inadimplentes_qs())
        self.assertFalse(self.receita.esta_atrasada)

    def test_indicadores_do_imovel_mostram_o_recebido(self):
        from patrimonio.indicadores import indicadores_do_imovel
        indicadores = indicadores_do_imovel(self.imovel, referencia=self.hoje)
        self.assertEqual(indicadores['receita_12m'], Decimal('500.00'))

    def test_imovel_detail_total_recebido_mostra_500(self):
        user = com_leitura(User.objects.create_user('caixa_cancel_user', password='pass'))
        client = Client()
        client.login(username='caixa_cancel_user', password='pass')
        resposta = client.get(reverse('imovel_detail', args=[self.imovel.pk]))
        self.assertEqual(resposta.context['total_recebido'], Decimal('500.00'))

    def test_dashboard_receitas_recebidas_mostra_500(self):
        user = com_leitura(User.objects.create_user('caixa_cancel_dash', password='pass'))
        client = Client()
        client.login(username='caixa_cancel_dash', password='pass')
        resposta = client.get(reverse('dashboard'))
        self.assertEqual(resposta.context['receitas_recebidas'], Decimal('500.00'))

    def test_resumo_por_imovel_dos_exports_mostra_500(self):
        from financeiro.exports import _resumo_por_imovel
        resumo = dict(_resumo_por_imovel([self.receita], []))
        self.assertEqual(resumo[self.imovel.nome]['rec_rec'], Decimal('500.00'))
        self.assertEqual(resumo[self.imovel.nome]['em_aberto'], 0)

    # ─── Item 8 (rodada pós-revisão): valor lançado × exigível × recebido ────

    def test_dashboard_receitas_previstas_exclui_cancelada(self):
        """Valor exigível (item 8): a receita cancelada não conta em 'Receita Prevista'."""
        user = com_leitura(User.objects.create_user('caixa_cancel_prev', password='pass'))
        client = Client()
        client.login(username='caixa_cancel_prev', password='pass')
        resposta = client.get(reverse('dashboard'))
        self.assertEqual(resposta.context['receitas_previstas'], 0)
        # o recebido continua contando (500), mesmo com previsto zerado
        self.assertEqual(resposta.context['receitas_recebidas'], Decimal('500.00'))

    def test_resumo_por_imovel_exigivel_exclui_cancelada_e_guarda_lancado(self):
        from financeiro.exports import _resumo_por_imovel
        resumo = dict(_resumo_por_imovel([self.receita], []))
        item = resumo[self.imovel.nome]
        self.assertEqual(item['rec_prev'], Decimal('0.00'))  # exigível: exclui cancelada
        self.assertEqual(item['rec_cancel'], Decimal('2000.00'))  # lançado histórico, à parte
        self.assertEqual(item['rec_rec'], Decimal('500.00'))  # recebido: sempre conta

    def test_relatorio_mensal_resumo_exclui_cancelada_da_receita_exigivel(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest('openpyxl não instalado')
        import io
        client = Client()
        com_leitura(User.objects.create_user('caixa_cancel_rel', password='pass'))
        client.login(username='caixa_cancel_rel', password='pass')
        response = client.get(
            reverse('export_relatorio_mensal', args=['xlsx']),
            {'mes': self.hoje.month, 'ano': self.hoje.year},
        )
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb['Resumo por Imóvel']
        linhas = list(ws.iter_rows(min_row=2, values_only=True))
        linha = next(l for l in linhas if l[0] == self.imovel.nome)
        # colunas: Imóvel, Receita Exigível, Receita Cancelada (Lançado), Receita Recebida, Despesas Pagas, Resultado
        self.assertEqual(linha[1], 0.0)
        self.assertEqual(linha[2], 2000.0)
        self.assertEqual(linha[3], 500.0)

    def test_relatorio_contabil_resumo_exclui_cancelada_da_receita_exigivel(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest('openpyxl não instalado')
        import io
        client = Client()
        com_leitura(User.objects.create_user('caixa_cancel_cont', password='pass'))
        client.login(username='caixa_cancel_cont', password='pass')
        response = client.get(
            reverse('export_relatorio_contabilidade'),
            {'mes': self.hoje.month, 'ano': self.hoje.year},
        )
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb['Resumo']
        campos = {row[0]: row[1] for row in ws.iter_rows(min_row=2, values_only=True) if row[0]}
        self.assertEqual(campos['Total Receitas Exigíveis (R$)'], 0.0)
        self.assertEqual(campos['Total Receitas Canceladas — Lançado (R$)'], 2000.0)
        self.assertEqual(campos['Total Receitas Recebidas (R$)'], 500.0)


class ValorExigivelReceitaCanceladaTest(TestCase):
    """
    Item 8: cenários adicionais de receita cancelada sem recebimento e com
    pagamento integral, confirmando que o valor exigível some (0) em ambos,
    mas o valor lançado e o recebido seguem as regras corretas.
    """

    def setUp(self):
        self.imovel, self.locatario = _criar_base()
        self.contrato = _criar_contrato(
            self.imovel, self.locatario, data_inicio=date(2020, 1, 1), data_fim=date(2030, 12, 31),
        )
        self.hoje = timezone.localdate()

    def test_cancelada_sem_recebimento_nao_conta_como_exigivel(self):
        from financeiro.exports import _resumo_por_imovel
        receita = ReceitaAluguel.objects.create(
            contrato=self.contrato, imovel=self.imovel,
            competencia_mes=self.hoje.month, competencia_ano=self.hoje.year,
            data_vencimento=self.hoje, valor_previsto=Decimal('1500.00'), status='previsto',
        )
        receita.cancelar()
        resumo = dict(_resumo_por_imovel([receita], []))
        item = resumo[self.imovel.nome]
        self.assertEqual(item['rec_prev'], Decimal('0.00'))
        self.assertEqual(item['rec_cancel'], Decimal('1500.00'))
        self.assertEqual(item['rec_rec'], Decimal('0.00'))

    def test_cancelada_apos_pagamento_integral_nao_conta_como_exigivel(self):
        from financeiro.exports import _resumo_por_imovel
        from financeiro.services import registrar_recebimento
        receita = ReceitaAluguel.objects.create(
            contrato=self.contrato, imovel=self.imovel,
            competencia_mes=self.hoje.month, competencia_ano=self.hoje.year,
            data_vencimento=self.hoje, valor_previsto=Decimal('1000.00'), status='previsto',
        )
        registrar_recebimento(receita.pk, Decimal('1000.00'), self.hoje)
        receita.refresh_from_db()
        receita.cancelar()
        resumo = dict(_resumo_por_imovel([receita], []))
        item = resumo[self.imovel.nome]
        self.assertEqual(item['rec_prev'], Decimal('0.00'))
        self.assertEqual(item['rec_cancel'], Decimal('1000.00'))
        self.assertEqual(item['rec_rec'], Decimal('1000.00'))
        self.assertEqual(receita.saldo_em_aberto, Decimal('0.00'))

    def test_receitas_list_total_exigivel_exclui_cancelada(self):
        outro_contrato = _criar_contrato(
            self.imovel, self.locatario, data_inicio=date(2020, 1, 1), data_fim=date(2030, 12, 31),
        )
        ReceitaAluguel.objects.create(
            contrato=self.contrato, imovel=self.imovel,
            competencia_mes=self.hoje.month, competencia_ano=self.hoje.year,
            data_vencimento=self.hoje, valor_previsto=Decimal('1000.00'), status='previsto',
        )
        cancelada = ReceitaAluguel.objects.create(
            contrato=outro_contrato, imovel=self.imovel,
            competencia_mes=self.hoje.month, competencia_ano=self.hoje.year,
            data_vencimento=self.hoje, valor_previsto=Decimal('500.00'), status='previsto',
        )
        cancelada.cancelar()

        client = Client()
        com_leitura(User.objects.create_user('receitas_list_total', password='pass'))
        client.login(username='receitas_list_total', password='pass')
        response = client.get(
            reverse('receitas_list'), {'mes': self.hoje.month, 'ano': self.hoje.year},
        )
        self.assertEqual(response.context['total_exigivel'], Decimal('1000.00'))
        conteudo = response.content.decode()
        self.assertIn('Total exigível', conteudo)


class ReceitaExigivelComEncargosTest(TestCase):
    """
    Item 7 (rodada final): "Receita Exigível"/"Total Exigível"/"Valor
    Exigível" é SEMPRE valor_previsto + multa + juros − desconto — nunca
    valor_previsto isolado. Exemplo do enunciado: previsto R$ 2.000, multa
    R$ 100, juros R$ 50, desconto R$ 20 → total devido R$ 2.130.
    """
    def setUp(self):
        self.imovel, self.locatario = _criar_base()
        self.contrato = _criar_contrato(
            self.imovel, self.locatario, data_inicio=date(2020, 1, 1), data_fim=date(2030, 12, 31),
        )
        self.hoje = timezone.localdate()

    def _receita(self, **kwargs):
        defaults = dict(
            contrato=self.contrato, imovel=self.imovel,
            competencia_mes=self.hoje.month, competencia_ano=self.hoje.year,
            data_vencimento=self.hoje, valor_previsto=Decimal('2000.00'), status='previsto',
        )
        defaults.update(kwargs)
        return ReceitaAluguel.objects.create(**defaults)

    def test_receita_simples_sem_encargos(self):
        r = self._receita()
        self.assertEqual(r.valor_total_devido, Decimal('2000.00'))

    def test_receita_com_multa(self):
        r = self._receita(multa=Decimal('100.00'))
        self.assertEqual(r.valor_total_devido, Decimal('2100.00'))

    def test_receita_com_juros(self):
        r = self._receita(juros=Decimal('50.00'))
        self.assertEqual(r.valor_total_devido, Decimal('2050.00'))

    def test_receita_com_desconto(self):
        r = self._receita(desconto=Decimal('20.00'))
        self.assertEqual(r.valor_total_devido, Decimal('1980.00'))

    def test_combinacao_multa_juros_desconto(self):
        r = self._receita(multa=Decimal('100.00'), juros=Decimal('50.00'), desconto=Decimal('20.00'))
        self.assertEqual(r.valor_total_devido, Decimal('2130.00'))

    def test_cancelada_com_encargos_nao_conta_no_exigivel_agregado(self):
        r = self._receita(multa=Decimal('100.00'), juros=Decimal('50.00'), desconto=Decimal('20.00'))
        r.cancelar()
        total = ReceitaAluguel.objects.exclude(status='cancelado').aggregate(
            total=models.Sum(valor_total_devido_expr())
        )['total']
        self.assertIsNone(total)

    def test_parcialmente_recebida_mantem_valor_total_devido_correto(self):
        from financeiro.services import registrar_recebimento
        r = self._receita(multa=Decimal('100.00'), juros=Decimal('50.00'), desconto=Decimal('20.00'))
        registrar_recebimento(r.pk, Decimal('500.00'), self.hoje)
        r.refresh_from_db()
        self.assertEqual(r.valor_total_devido, Decimal('2130.00'))
        self.assertEqual(r.saldo_em_aberto, Decimal('1630.00'))

    def test_dashboard_soma_valor_total_devido_nao_valor_previsto_isolado(self):
        self._receita(multa=Decimal('100.00'), juros=Decimal('50.00'), desconto=Decimal('20.00'))
        client = Client()
        com_leitura(User.objects.create_user('exig_dash', password='pass'))
        client.login(username='exig_dash', password='pass')
        resposta = client.get(reverse('receitas_list'), {'mes': self.hoje.month, 'ano': self.hoje.year})
        self.assertEqual(resposta.context['total_exigivel'], Decimal('2130.00'))

    def test_resumo_por_imovel_usa_valor_total_devido(self):
        from financeiro.exports import _resumo_por_imovel
        r = self._receita(multa=Decimal('100.00'), juros=Decimal('50.00'), desconto=Decimal('20.00'))
        resumo = dict(_resumo_por_imovel([r], []))
        self.assertEqual(resumo[self.imovel.nome]['rec_prev'], Decimal('2130.00'))

    def test_relatorio_mensal_exporta_valor_total_devido(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest('openpyxl não instalado')
        import io
        self._receita(multa=Decimal('100.00'), juros=Decimal('50.00'), desconto=Decimal('20.00'))
        client = Client()
        com_leitura(User.objects.create_user('exig_rel_mensal', password='pass'))
        client.login(username='exig_rel_mensal', password='pass')
        response = client.get(
            reverse('export_relatorio_mensal', args=['xlsx']),
            {'mes': self.hoje.month, 'ano': self.hoje.year},
        )
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb['Resumo por Imóvel']
        linha = next(l for l in ws.iter_rows(min_row=2, values_only=True) if l[0] == self.imovel.nome)
        self.assertEqual(linha[1], 2130.0)

    def test_relatorio_contabil_exporta_valor_total_devido(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest('openpyxl não instalado')
        import io
        self._receita(multa=Decimal('100.00'), juros=Decimal('50.00'), desconto=Decimal('20.00'))
        client = Client()
        com_leitura(User.objects.create_user('exig_rel_cont', password='pass'))
        client.login(username='exig_rel_cont', password='pass')
        response = client.get(
            reverse('export_relatorio_contabilidade'),
            {'mes': self.hoje.month, 'ano': self.hoje.year},
        )
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws_resumo = wb['Resumo']
        campos = {row[0]: row[1] for row in ws_resumo.iter_rows(min_row=2, values_only=True) if row[0]}
        self.assertEqual(campos['Total Receitas Exigíveis (R$)'], 2130.0)
        ws_por_im = wb['Por Imóvel']
        linha = next(l for l in ws_por_im.iter_rows(min_row=2, values_only=True) if l[0] == self.imovel.nome)
        self.assertEqual(linha[1], 2130.0)

    def test_export_receitas_csv_e_xlsx_distinguem_previsto_de_total_devido(self):
        r = self._receita(multa=Decimal('100.00'), juros=Decimal('50.00'), desconto=Decimal('20.00'))
        client = Client()
        com_leitura(User.objects.create_user('exig_csv_xlsx', password='pass'))
        client.login(username='exig_csv_xlsx', password='pass')

        resp_csv = client.get(
            reverse('export_receitas', args=['csv']), {'mes': self.hoje.month, 'ano': self.hoje.year},
        )
        conteudo_csv = resp_csv.content.decode('utf-8-sig')
        self.assertIn('2000', conteudo_csv)  # Valor Previsto
        self.assertIn('2130', conteudo_csv)  # Valor Total Devido

        try:
            import openpyxl
        except ImportError:
            self.skipTest('openpyxl não instalado')
        import io
        resp_xlsx = client.get(
            reverse('export_receitas', args=['xlsx']), {'mes': self.hoje.month, 'ano': self.hoje.year},
        )
        wb = openpyxl.load_workbook(io.BytesIO(resp_xlsx.content))
        ws = wb['Receitas']
        cabecalho = [c.value for c in ws[1]]
        linha = dict(zip(cabecalho, next(ws.iter_rows(min_row=2, max_row=2, values_only=True))))
        self.assertEqual(linha['Valor Previsto'], 2000.0)
        self.assertEqual(linha['Valor Total Devido'], 2130.0)


# ─── Item 10 (rodada pós-revisão): integridade do model Despesa ───────────────

class DespesaIntegridadeTest(TestCase):
    def setUp(self):
        self.imovel, _ = _criar_base()
        self.hoje = timezone.localdate()

    def _despesa(self, **kwargs):
        defaults = dict(
            imovel=self.imovel, descricao='Despesa teste', categoria='outro',
            data_vencimento=self.hoje, valor=Decimal('100.00'), status='prevista',
        )
        defaults.update(kwargs)
        return Despesa(**defaults)

    def test_valor_zero_rejeitado(self):
        despesa = self._despesa(valor=Decimal('0.00'))
        with self.assertRaises(ValidationError):
            despesa.full_clean()

    def test_valor_negativo_rejeitado(self):
        despesa = self._despesa(valor=Decimal('-50.00'))
        with self.assertRaises(ValidationError):
            despesa.full_clean()

    def test_paga_sem_data_rejeitada(self):
        despesa = self._despesa(status='paga', data_pagamento=None)
        with self.assertRaises(ValidationError) as ctx:
            despesa.full_clean()
        self.assertIn('data_pagamento', ctx.exception.message_dict)

    def test_prevista_com_data_rejeitada(self):
        despesa = self._despesa(status='prevista', data_pagamento=self.hoje)
        with self.assertRaises(ValidationError) as ctx:
            despesa.full_clean()
        self.assertIn('data_pagamento', ctx.exception.message_dict)

    def test_atrasada_com_data_rejeitada(self):
        despesa = self._despesa(status='atrasada', data_pagamento=self.hoje)
        with self.assertRaises(ValidationError) as ctx:
            despesa.full_clean()
        self.assertIn('data_pagamento', ctx.exception.message_dict)

    def test_cancelada_com_data_rejeitada(self):
        despesa = self._despesa(status='cancelada', data_pagamento=self.hoje)
        with self.assertRaises(ValidationError) as ctx:
            despesa.full_clean()
        self.assertIn('data_pagamento', ctx.exception.message_dict)

    def test_valor_zero_rejeitado_no_banco_mesmo_sem_full_clean(self):
        """A constraint do banco protege mesmo criação direta via ORM sem full_clean()."""
        from django.db import IntegrityError, transaction as db_transaction
        with self.assertRaises(IntegrityError):
            with db_transaction.atomic():
                Despesa.objects.create(
                    imovel=self.imovel, descricao='Sem clean', categoria='outro',
                    data_vencimento=self.hoje, valor=Decimal('0.00'), status='prevista',
                )

    def test_status_data_incoerentes_rejeitados_no_banco_mesmo_sem_full_clean(self):
        from django.db import IntegrityError, transaction as db_transaction
        with self.assertRaises(IntegrityError):
            with db_transaction.atomic():
                Despesa.objects.create(
                    imovel=self.imovel, descricao='Incoerente', categoria='outro',
                    data_vencimento=self.hoje, valor=Decimal('100.00'),
                    status='paga', data_pagamento=None,
                )

    def test_marcar_como_paga(self):
        from financeiro.services import marcar_despesa_como_paga
        despesa = Despesa.objects.create(
            imovel=self.imovel, descricao='A pagar', categoria='outro',
            data_vencimento=self.hoje, valor=Decimal('250.00'), status='prevista',
        )
        marcar_despesa_como_paga(despesa.pk, self.hoje)
        despesa.refresh_from_db()
        self.assertEqual(despesa.status, 'paga')
        self.assertEqual(despesa.data_pagamento, self.hoje)

    def test_marcar_como_paga_sem_data_e_rejeitado(self):
        from financeiro.services import marcar_despesa_como_paga
        despesa = Despesa.objects.create(
            imovel=self.imovel, descricao='A pagar 2', categoria='outro',
            data_vencimento=self.hoje, valor=Decimal('250.00'), status='prevista',
        )
        with self.assertRaises(ValidationError):
            marcar_despesa_como_paga(despesa.pk, None)

    def test_reabrir_despesa_limpa_data(self):
        from financeiro.services import marcar_despesa_como_paga, reabrir_despesa
        despesa = Despesa.objects.create(
            imovel=self.imovel, descricao='A reabrir', categoria='outro',
            data_vencimento=self.hoje + timedelta(days=10), valor=Decimal('300.00'), status='prevista',
        )
        marcar_despesa_como_paga(despesa.pk, self.hoje)
        reabrir_despesa(despesa.pk)
        despesa.refresh_from_db()
        self.assertIsNone(despesa.data_pagamento)
        self.assertEqual(despesa.status, 'prevista')

    def test_reabrir_despesa_vencida_marca_atrasada(self):
        from financeiro.services import marcar_despesa_como_paga, reabrir_despesa
        despesa = Despesa.objects.create(
            imovel=self.imovel, descricao='A reabrir vencida', categoria='outro',
            data_vencimento=self.hoje - timedelta(days=5), valor=Decimal('300.00'), status='prevista',
        )
        marcar_despesa_como_paga(despesa.pk, self.hoje)
        reabrir_despesa(despesa.pk)
        despesa.refresh_from_db()
        self.assertIsNone(despesa.data_pagamento)
        self.assertEqual(despesa.status, 'atrasada')

    def test_cancelar_despesa_limpa_data(self):
        from financeiro.services import marcar_despesa_como_paga, cancelar_despesa
        despesa = Despesa.objects.create(
            imovel=self.imovel, descricao='A cancelar', categoria='outro',
            data_vencimento=self.hoje, valor=Decimal('300.00'), status='prevista',
        )
        marcar_despesa_como_paga(despesa.pk, self.hoje)
        cancelar_despesa(despesa.pk)
        despesa.refresh_from_db()
        self.assertIsNone(despesa.data_pagamento)
        self.assertEqual(despesa.status, 'cancelada')

    def test_conciliacao_de_debito_cria_estado_coerente(self):
        from conciliacao.models import ContaBancaria, ExtratoImportado, TransacaoExtrato
        from conciliacao.services import lancar_despesa
        conta = ContaBancaria.objects.create(nome='Conta Débito')
        extrato = ExtratoImportado.objects.create(conta=conta, hash_arquivo='hdesp10')
        transacao = TransacaoExtrato.objects.create(
            extrato=extrato, conta=conta, fitid='td1', data=self.hoje,
            valor=Decimal('-150.00'), tipo='debito', descricao='CEMIG',
        )
        despesa = lancar_despesa(transacao, categoria='outro', descricao='Energia', fornecedor=None, imovel=None)
        self.assertEqual(despesa.status, 'paga')
        self.assertEqual(despesa.data_pagamento, self.hoje)
        despesa.full_clean()  # não deve levantar — estado coerente

    def test_desfazer_repasse_reabre_comissao_e_limpa_data(self):
        from patrimonio.models import Pessoa, Contrato
        from conciliacao.models import ContaBancaria, ExtratoImportado, TransacaoExtrato
        from conciliacao.services import conciliar_com_receitas, desfazer_conciliacao
        from financeiro.services import gerar_receitas_para_contrato

        locatario = Pessoa.objects.create(nome='Locatário Desp10', tipo='locatario')
        imob = Pessoa.objects.create(nome='Imob Desp10', tipo='imobiliaria')
        contrato = Contrato.objects.create(
            imovel=self.imovel, locatario=locatario, imobiliaria=imob,
            data_inicio=date(2026, 1, 1), data_fim=date(2026, 12, 31),
            dia_vencimento=10, valor_aluguel=Decimal('2000.00'),
            comissao_imobiliaria_percentual=Decimal('10.00'),
        )
        gerar_receitas_para_contrato(contrato, data_inicio=date(2026, 3, 1), data_fim=date(2026, 3, 31))
        receita = ReceitaAluguel.objects.get(contrato=contrato)
        comissao = Despesa.objects.get(contrato=contrato, origem_automatica=True)

        conta = ContaBancaria.objects.create(nome='Conta Repasse10')
        extrato = ExtratoImportado.objects.create(conta=conta, hash_arquivo='hdesp10b')
        transacao = TransacaoExtrato.objects.create(
            extrato=extrato, conta=conta, fitid='td2', data=date(2026, 3, 12),
            valor=Decimal('1800.00'), tipo='credito', descricao='REPASSE',
        )
        conciliar_com_receitas(transacao, [receita], marcar_comissoes=True)
        comissao.refresh_from_db()
        self.assertEqual(comissao.status, 'paga')
        self.assertIsNotNone(comissao.data_pagamento)

        desfazer_conciliacao(transacao)
        comissao.refresh_from_db()
        self.assertEqual(comissao.status, 'prevista')
        self.assertIsNone(comissao.data_pagamento)
        comissao.full_clean()  # estado coerente após desfazer
