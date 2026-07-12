import csv
from decimal import Decimal

from django.http import HttpResponse
from django.utils import timezone

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

from .models import ReceitaAluguel, Despesa, receitas_inadimplentes_qs


# ─── helpers ───────────────────────────────────────────────────────────────────

def _estilizar_cabecalho(ws, num_colunas):
    """Aplica estilo azul escuro nas células do cabeçalho (linha 1)."""
    fill = PatternFill(start_color='1F3864', end_color='1F3864', fill_type='solid')
    font = Font(color='FFFFFF', bold=True)
    for col in range(1, num_colunas + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal='center')
    for col in range(1, num_colunas + 1):
        ws.column_dimensions[get_column_letter(col)].width = 20


def _response_xlsx(nome_arquivo):
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{nome_arquivo}"'
    return response


def _response_csv(nome_arquivo):
    response = HttpResponse(content_type='text/csv; charset=utf-8-sig')
    response['Content-Disposition'] = f'attachment; filename="{nome_arquivo}"'
    return response


# ─── imóveis ───────────────────────────────────────────────────────────────────

CABECALHO_IMOVEIS = [
    'ID', 'Nome', 'Endereço', 'Cidade', 'Estado', 'CEP',
    'Matrícula', 'Inscrição IPTU', 'Proprietário', 'Status',
    'Tipo', 'Uso', 'Imóvel Pai',
    'Data Aquisição', 'Valor Aquisição', 'Valor Estimado',
]


def _linhas_imoveis(imoveis):
    for i in imoveis:
        yield [
            i.pk, i.nome, i.endereco, i.cidade, i.estado, i.cep,
            i.matricula, i.inscricao_iptu, i.proprietario,
            i.get_status_display(),
            i.get_tipo_imovel_display(), i.get_uso_display(),
            i.imovel_pai.nome if i.imovel_pai else '',
            i.data_aquisicao, i.valor_aquisicao, i.valor_estimado,
        ]


def exportar_imoveis_csv(request, imoveis):
    response = _response_csv('imoveis.csv')
    writer = csv.writer(response)
    writer.writerow(CABECALHO_IMOVEIS)
    for linha in _linhas_imoveis(imoveis):
        writer.writerow(linha)
    return response


def exportar_imoveis_xlsx(request, imoveis):
    if not OPENPYXL_AVAILABLE:
        return exportar_imoveis_csv(request, imoveis)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Imóveis'
    ws.append(CABECALHO_IMOVEIS)
    for linha in _linhas_imoveis(imoveis):
        ws.append(linha)
    _estilizar_cabecalho(ws, len(CABECALHO_IMOVEIS))
    response = _response_xlsx('imoveis.xlsx')
    wb.save(response)
    return response


# ─── contratos ─────────────────────────────────────────────────────────────────

CABECALHO_CONTRATOS = [
    'ID', 'Imóvel', 'Locatário(s)', 'Fiador(es)', 'Imobiliária',
    'Data Início', 'Data Fim', 'Prazo Indeterminado',
    'Valor Aluguel', 'Dia Vencimento', 'Índice', 'Próximo Reajuste',
    'Garantia', 'Status',
]


def _linhas_contratos(contratos):
    for c in contratos:
        imobiliaria = c.get_imobiliaria_principal()
        yield [
            c.pk, c.imovel.nome, c.locatarios_display, c.fiadores_display,
            imobiliaria.nome if imobiliaria else '',
            c.data_inicio, c.data_fim, 'Sim' if c.prazo_indeterminado else 'Não',
            c.valor_aluguel,
            c.dia_vencimento, c.get_indice_reajuste_display(),
            c.data_proximo_reajuste, c.get_tipo_garantia_display(),
            c.get_status_display(),
        ]


def exportar_contratos_csv(request, contratos):
    response = _response_csv('contratos.csv')
    writer = csv.writer(response)
    writer.writerow(CABECALHO_CONTRATOS)
    for linha in _linhas_contratos(contratos):
        writer.writerow(linha)
    return response


def exportar_contratos_xlsx(request, contratos):
    if not OPENPYXL_AVAILABLE:
        return exportar_contratos_csv(request, contratos)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Contratos'
    ws.append(CABECALHO_CONTRATOS)
    for linha in _linhas_contratos(contratos):
        ws.append(linha)
    _estilizar_cabecalho(ws, len(CABECALHO_CONTRATOS))
    response = _response_xlsx('contratos.xlsx')
    wb.save(response)
    return response


# ─── receitas ──────────────────────────────────────────────────────────────────

CABECALHO_RECEITAS = [
    'ID', 'Imóvel', 'Locatário', 'Mês', 'Ano',
    'Vencimento', 'Valor Previsto', 'Multa', 'Juros', 'Desconto',
    'Valor Total Devido', 'Valor Recebido', 'Saldo em Aberto',
    'Data Recebimento', 'Status',
]


def _linhas_receitas(receitas):
    for r in receitas:
        locatario = r.contrato.locatarios_display if r.contrato else ''
        yield [
            r.pk, r.imovel.nome, locatario,
            r.get_competencia_mes_display(), r.competencia_ano,
            r.data_vencimento, r.valor_previsto, r.multa, r.juros, r.desconto,
            r.valor_total_devido, r.valor_recebido, r.saldo_em_aberto,
            r.data_recebimento, r.get_status_display(),
        ]


def exportar_receitas_csv(request, receitas):
    response = _response_csv('receitas.csv')
    writer = csv.writer(response)
    writer.writerow(CABECALHO_RECEITAS)
    for linha in _linhas_receitas(receitas):
        writer.writerow(linha)
    return response


def exportar_receitas_xlsx(request, receitas):
    if not OPENPYXL_AVAILABLE:
        return exportar_receitas_csv(request, receitas)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Receitas'
    ws.append(CABECALHO_RECEITAS)
    for linha in _linhas_receitas(receitas):
        ws.append(linha)
    _estilizar_cabecalho(ws, len(CABECALHO_RECEITAS))
    response = _response_xlsx('receitas.xlsx')
    wb.save(response)
    return response


# ─── despesas ──────────────────────────────────────────────────────────────────

CABECALHO_DESPESAS = [
    'ID', 'Imóvel', 'Categoria', 'Fornecedor', 'Descrição',
    'Mês', 'Ano', 'Vencimento', 'Valor', 'Data Pagamento', 'Status',
]


def _linhas_despesas(despesas):
    for d in despesas:
        yield [
            d.pk,
            d.imovel.nome if d.imovel else '',
            d.get_categoria_display(),
            d.fornecedor.nome if d.fornecedor else '',
            d.descricao,
            d.get_competencia_mes_display() if d.competencia_mes else '',
            d.competencia_ano or '',
            d.data_vencimento, d.valor, d.data_pagamento,
            d.get_status_display(),
        ]


def exportar_despesas_csv(request, despesas):
    response = _response_csv('despesas.csv')
    writer = csv.writer(response)
    writer.writerow(CABECALHO_DESPESAS)
    for linha in _linhas_despesas(despesas):
        writer.writerow(linha)
    return response


def exportar_despesas_xlsx(request, despesas):
    if not OPENPYXL_AVAILABLE:
        return exportar_despesas_csv(request, despesas)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Despesas'
    ws.append(CABECALHO_DESPESAS)
    for linha in _linhas_despesas(despesas):
        ws.append(linha)
    _estilizar_cabecalho(ws, len(CABECALHO_DESPESAS))
    response = _response_xlsx('despesas.xlsx')
    wb.save(response)
    return response


# ─── inadimplência ─────────────────────────────────────────────────────────────

CABECALHO_INADIMPLENCIA = [
    'Imóvel', 'Locatário', 'Mês', 'Ano', 'Vencimento',
    'Valor Total Devido', 'Valor Recebido', 'Saldo em Aberto', 'Dias Atraso',
]


def _linhas_inadimplencia(receitas_atrasadas):
    hoje = timezone.localdate()
    for r in receitas_atrasadas:
        dias = (hoje - r.data_vencimento).days
        locatario = r.contrato.locatarios_display if r.contrato else ''
        yield [
            r.imovel.nome, locatario,
            r.get_competencia_mes_display(), r.competencia_ano,
            r.data_vencimento, r.valor_total_devido, r.valor_recebido or 0,
            r.saldo_em_aberto, dias,
        ]


def exportar_inadimplencia_csv(request, receitas_atrasadas):
    response = _response_csv('inadimplencia.csv')
    writer = csv.writer(response)
    writer.writerow(CABECALHO_INADIMPLENCIA)
    for linha in _linhas_inadimplencia(receitas_atrasadas):
        writer.writerow(linha)
    return response


def exportar_inadimplencia_xlsx(request, receitas_atrasadas):
    if not OPENPYXL_AVAILABLE:
        return exportar_inadimplencia_csv(request, receitas_atrasadas)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Inadimplência'
    ws.append(CABECALHO_INADIMPLENCIA)
    for linha in _linhas_inadimplencia(receitas_atrasadas):
        ws.append(linha)
    _estilizar_cabecalho(ws, len(CABECALHO_INADIMPLENCIA))
    response = _response_xlsx('inadimplencia.xlsx')
    wb.save(response)
    return response


# ─── resumo por imóvel (agrupamento em memória, uma passada) ───────────────────

def _resumo_por_imovel(receitas, despesas):
    """
    Agrupa receitas/despesas por imóvel em uma única passada.
    Retorna lista ordenada por nome: [(nome, {rec_prev, rec_rec, rec_cancel, desp_pagas, em_aberto}), ...],
    omitindo imóveis sem nenhum valor no período.

    rec_prev (valor EXIGÍVEL — item 8): soma valor_previsto só das receitas
    NÃO canceladas — a cobrança de uma cancelada foi encerrada, não é mais
    algo a exigir. rec_cancel guarda separadamente o valor_previsto LANÇADO
    das canceladas (histórico), para quem precisar do total original.
    """
    from collections import defaultdict
    from decimal import Decimal

    dados = defaultdict(lambda: {
        'rec_prev': Decimal('0.00'), 'rec_rec': Decimal('0.00'),
        'rec_cancel': Decimal('0.00'), 'desp_pagas': Decimal('0.00'), 'em_aberto': 0,
    })
    for r in receitas:
        item = dados[r.imovel.nome]
        if r.status == 'cancelado':
            item['rec_cancel'] += r.valor_previsto
        else:
            item['rec_prev'] += r.valor_previsto
        # Sempre soma valor_recebido, mesmo cancelada: recalcular_recebimentos
        # reconsolida esse campo a partir dos RecebimentoReceita reais
        # independentemente do status — dinheiro recebido não desaparece com
        # o cancelamento (item 4).
        item['rec_rec'] += r.valor_recebido or Decimal('0.00')
        # Em aberto pelo SALDO real (cancelada tem saldo 0) — não pelo status
        if r.saldo_em_aberto > 0:
            item['em_aberto'] += 1
    for d in despesas:
        if d.imovel_id and d.status == 'paga':
            dados[d.imovel.nome]['desp_pagas'] += d.valor

    return sorted(
        (nome, item) for nome, item in dados.items()
        if item['rec_prev'] or item['rec_rec'] or item['rec_cancel'] or item['desp_pagas']
    )


# ─── relatório mensal ──────────────────────────────────────────────────────────

def exportar_relatorio_mensal_xlsx(request, mes, ano):
    if not OPENPYXL_AVAILABLE:
        return HttpResponse('openpyxl não instalado.', status=500)

    wb = openpyxl.Workbook()

    # Aba receitas
    ws_rec = wb.active
    ws_rec.title = 'Receitas'
    ws_rec.append(CABECALHO_RECEITAS)
    receitas = ReceitaAluguel.objects.filter(
        competencia_mes=mes, competencia_ano=ano
    ).select_related('imovel', 'contrato__locatario')
    for linha in _linhas_receitas(receitas):
        ws_rec.append(linha)
    _estilizar_cabecalho(ws_rec, len(CABECALHO_RECEITAS))

    # Aba despesas
    ws_desp = wb.create_sheet('Despesas')
    ws_desp.append(CABECALHO_DESPESAS)
    despesas = Despesa.objects.filter(
        competencia_mes=mes, competencia_ano=ano
    ).select_related('imovel', 'fornecedor')
    for linha in _linhas_despesas(despesas):
        ws_desp.append(linha)
    _estilizar_cabecalho(ws_desp, len(CABECALHO_DESPESAS))

    # Aba resumo por imóvel — agrupamento em uma passada só
    ws_res = wb.create_sheet('Resumo por Imóvel')
    ws_res.append([
        'Imóvel', 'Receita Exigível', 'Receita Cancelada (Lançado)',
        'Receita Recebida', 'Despesas Pagas', 'Resultado',
    ])
    _estilizar_cabecalho(ws_res, 6)
    resumo = _resumo_por_imovel(receitas, despesas)
    for nome, dados in resumo:
        ws_res.append([
            nome, dados['rec_prev'], dados['rec_cancel'], dados['rec_rec'], dados['desp_pagas'],
            dados['rec_rec'] - dados['desp_pagas'],
        ])

    response = _response_xlsx(f'relatorio_{mes:02d}_{ano}.xlsx')
    wb.save(response)
    return response


# ─── relatório contabilidade ───────────────────────────────────────────────────

def exportar_relatorio_contabilidade_xlsx(request, mes, ano):
    if not OPENPYXL_AVAILABLE:
        return HttpResponse('openpyxl não instalado.', status=500)

    from documentos.models import Documento

    # timezone.localdate(): data de negócio no fuso do Django (America/Sao_Paulo)
    # — timezone.now().date() seria a data UTC e viraria o dia às 21h locais.
    hoje = timezone.localdate()
    receitas = list(
        ReceitaAluguel.objects.filter(competencia_mes=mes, competencia_ano=ano)
        .select_related('imovel', 'contrato__locatario')
    )
    despesas = list(
        Despesa.objects.filter(competencia_mes=mes, competencia_ano=ano)
        .select_related('imovel', 'fornecedor')
    )
    inadimplentes = list(
        receitas_inadimplentes_qs()
        .select_related('imovel', 'contrato__locatario')
        .order_by('data_vencimento')
    )
    docs_pendentes = list(
        Documento.objects.filter(
            tipo__in=Documento.TIPOS_CONTABILIDADE, enviado_contabilidade=False
        ).select_related('imovel', 'pessoa')
    )

    wb = openpyxl.Workbook()

    # ── Aba 1: Resumo ──────────────────────────────────────────────────────────
    ws_res = wb.active
    ws_res.title = 'Resumo'
    ws_res.append(['Campo', 'Valor'])
    _estilizar_cabecalho(ws_res, 2)
    # Valor EXIGÍVEL (item 8): exclui canceladas — a cobrança delas foi
    # encerrada. O valor LANÇADO das canceladas (histórico original) é
    # reportado à parte, nunca somado ao exigível.
    total_prev = sum(r.valor_previsto for r in receitas if r.status != 'cancelado')
    total_cancelado = sum(r.valor_previsto for r in receitas if r.status == 'cancelado')
    # Soma valor_recebido de TODAS as receitas, mesmo canceladas — dinheiro
    # recebido antes do cancelamento continua no caixa (item 4).
    total_rec = sum(r.valor_recebido or 0 for r in receitas)
    total_desp = sum(d.valor for d in despesas if d.status == 'paga')
    # Identificação pelo SALDO real (regra centralizada), não pelo status:
    # cancelada tem saldo 0; parcial/'recebido' com saldo > 0 conta como aberta.
    abertas = [r for r in receitas if r.saldo_em_aberto > 0]
    vencidas = [r for r in abertas if r.data_vencimento < hoje]
    rec_aberto = len(abertas)
    rec_vencidas = len(vencidas)
    saldo_aberto = sum((r.saldo_em_aberto for r in abertas), Decimal('0.00'))
    saldo_vencido = sum((r.saldo_em_aberto for r in vencidas), Decimal('0.00'))
    docs_pend_cnt = len(docs_pendentes)
    nome_mes = dict(ReceitaAluguel.MESES).get(mes, str(mes))
    for row in [
        ('Mês/Ano', f'{nome_mes}/{ano}'),
        ('Total Receitas Exigíveis (R$)', float(total_prev)),
        ('Total Receitas Canceladas — Lançado (R$)', float(total_cancelado)),
        ('Total Receitas Recebidas (R$)', float(total_rec)),
        ('Total Despesas Pagas (R$)', float(total_desp)),
        ('Resultado Líquido (R$)', float(total_rec - total_desp)),
        ('Receitas em Aberto', rec_aberto),
        ('Receitas Vencidas (inadimplência)', rec_vencidas),
        ('Saldo Total em Aberto (R$)', float(saldo_aberto)),
        ('Saldo Total Vencido (R$)', float(saldo_vencido)),
        ('Documentos Pendentes para Contabilidade', docs_pend_cnt),
        ('', ''),
        ('Nota — Inadimplência Aberta', 'A aba "Inadimplência Aberta" lista todas as receitas vencidas com saldo em aberto, independente do mês de competência.'),
    ]:
        ws_res.append(row)
    ws_res.column_dimensions['A'].width = 40
    ws_res.column_dimensions['B'].width = 25

    # ── Aba 2: Receitas ────────────────────────────────────────────────────────
    ws_rec = wb.create_sheet('Receitas')
    cab_rec = [
        'Imóvel', 'Locatário', 'Competência', 'Vencimento',
        'Valor Previsto', 'Valor Total Devido', 'Valor Recebido', 'Saldo em Aberto',
        'Data Recebimento', 'Status', 'Atrasada', 'Observações',
    ]
    ws_rec.append(cab_rec)
    _estilizar_cabecalho(ws_rec, len(cab_rec))
    for r in receitas:
        loc = r.contrato.locatarios_display if r.contrato else ''
        ws_rec.append([
            r.imovel.nome, loc,
            f'{r.get_competencia_mes_display()}/{r.competencia_ano}',
            r.data_vencimento, float(r.valor_previsto), float(r.valor_total_devido),
            float(r.valor_recebido) if r.valor_recebido else '',
            float(r.saldo_em_aberto),
            r.data_recebimento, r.get_status_display(),
            'Sim' if r.esta_atrasada else 'Não',
            r.observacoes,
        ])

    # ── Aba 3: Despesas ────────────────────────────────────────────────────────
    ws_desp = wb.create_sheet('Despesas')
    cab_desp = [
        'Imóvel', 'Categoria', 'Fornecedor', 'Descrição',
        'Competência', 'Vencimento', 'Valor', 'Data Pagamento',
        'Status', 'Atrasada', 'Observações',
    ]
    ws_desp.append(cab_desp)
    _estilizar_cabecalho(ws_desp, len(cab_desp))
    for d in despesas:
        comp = f'{d.get_competencia_mes_display()}/{d.competencia_ano}' if d.competencia_mes else ''
        ws_desp.append([
            d.imovel.nome if d.imovel else '',
            d.get_categoria_display(),
            d.fornecedor.nome if d.fornecedor else '',
            d.descricao, comp,
            d.data_vencimento, float(d.valor), d.data_pagamento,
            d.get_status_display(),
            'Sim' if d.esta_atrasada else 'Não',
            d.observacoes,
        ])

    # ── Aba 4: Inadimplência Aberta ────────────────────────────────────────────
    ws_inad = wb.create_sheet('Inadimplência Aberta')
    cab_inad = [
        'Imóvel', 'Locatário', 'Competência', 'Vencimento',
        'Valor Total Devido', 'Saldo em Aberto', 'Dias de Atraso', 'Status', 'Observações',
    ]
    ws_inad.append(cab_inad)
    _estilizar_cabecalho(ws_inad, len(cab_inad))
    for r in inadimplentes:
        loc = r.contrato.locatarios_display if r.contrato else ''
        dias = (hoje - r.data_vencimento).days
        ws_inad.append([
            r.imovel.nome, loc,
            f'{r.get_competencia_mes_display()}/{r.competencia_ano}',
            r.data_vencimento, float(r.valor_total_devido), float(r.saldo_em_aberto), dias,
            r.get_status_display(), r.observacoes,
        ])

    # ── Aba 5: Documentos Pendentes ────────────────────────────────────────────
    ws_docs = wb.create_sheet('Docs. Pendentes')
    cab_docs = ['Título', 'Tipo', 'Imóvel', 'Pessoa', 'Data Documento', 'Data Validade', 'Observações']
    ws_docs.append(cab_docs)
    _estilizar_cabecalho(ws_docs, len(cab_docs))
    for doc in docs_pendentes:
        ws_docs.append([
            doc.titulo, doc.get_tipo_display(),
            doc.imovel.nome if doc.imovel else '',
            doc.pessoa.nome if doc.pessoa else '',
            doc.data_documento, doc.data_validade, doc.observacoes,
        ])

    # ── Aba 6: Resultado por Imóvel ────────────────────────────────────────────
    ws_por_im = wb.create_sheet('Por Imóvel')
    cab_im = [
        'Imóvel', 'Rec. Exigível', 'Rec. Cancelada (Lançado)', 'Rec. Recebida',
        'Desp. Pagas', 'Resultado', 'Em Aberto',
    ]
    ws_por_im.append(cab_im)
    _estilizar_cabecalho(ws_por_im, len(cab_im))
    for nome, dados in _resumo_por_imovel(receitas, despesas):
        ws_por_im.append([
            nome, float(dados['rec_prev']), float(dados['rec_cancel']), float(dados['rec_rec']),
            float(dados['desp_pagas']), float(dados['rec_rec'] - dados['desp_pagas']),
            dados['em_aberto'],
        ])

    response = _response_xlsx(f'contabilidade_{mes:02d}_{ano}.xlsx')
    wb.save(response)
    return response
